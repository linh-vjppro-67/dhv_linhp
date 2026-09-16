from __future__ import annotations
from backend.document_upload import read_document_upload
from backend.email_templates import monthly_reminder, monthly_report, incoming_email, template_html
import hashlib,json,os,uuid,io,shutil,re,unicodedata,smtplib,html,base64,zipfile
from html.parser import HTMLParser
from urllib.parse import urlencode
from xml.etree import ElementTree as ET
from pathlib import Path
from dotenv import load_dotenv
load_dotenv(Path(__file__).resolve().parent.parent/'.env')
from email.message import EmailMessage
from datetime import date,datetime,timedelta
from typing import Optional,Literal
import pymupdf
import torch
import pytesseract
from PIL import Image
from docx import Document as WordDocument
from fastapi import Depends,FastAPI,File,Form,HTTPException,UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse,Response
from fastapi.staticfiles import StaticFiles
from fastapi.security import OAuth2PasswordBearer
from starlette.concurrency import run_in_threadpool
from app.qwen_chat import extract_document_title
from jose import JWTError,jwt
from passlib.context import CryptContext
from pydantic import BaseModel,Field
from sqlalchemy import Boolean,Date,DateTime,Float,ForeignKey,Integer,String,Text,UniqueConstraint,create_engine,or_,update
from sqlalchemy.orm import DeclarativeBase,Mapped,Session,mapped_column,sessionmaker
from app.qwen_chat import chat_search,generate_grounded_report,map_query_to_documents,status as qwen_status
from app.search_engine import search as ai_document_search
from app.indexer import sync_folder as ai_sync_folder
from app.db import fetch_chunks_for_documents,get_document_by_id as get_ai_document,get_indexed_documents
ROOT=Path(__file__).parent;STORE=ROOT/'storage';STORE.mkdir(exist_ok=True);ARCHIVE_ROOT=ROOT/'archive';ARCHIVE_ROOT.mkdir(exist_ok=True);AI_DOCUMENTS_ROOT=ROOT.parent/'documents'
REGISTER_TEMPLATE=ROOT/'templates'/'document-register-template.xlsx'
DOCUMENT_TYPE_CODES={'QUYẾT ĐỊNH': 'QĐ', 'THÔNG BÁO': 'TB', 'KẾ HOẠCH': 'KH', 'BÁO CÁO': 'BC', 'TỜ TRÌNH': 'TTr', 'HỢP ĐỒNG': 'HĐ', 'THOẢ THUẬN': 'TT', 'BIÊN BẢN': 'BB', 'NGHỊ QUYẾT HỘI ĐỒNG TRƯỜNG': 'NQ-HĐT', 'QUYẾT ĐỊNH HỘI ĐỒNG TRƯỜNG': 'QĐ-HĐT', 'GIẤY UỶ QUYỀN': 'GUQ', 'GIẤY XÁC NHẬN': 'GXN', 'GIẤY GIỚI THIỆU': 'GGT', 'GIẤY CHỨNG NHẬN': 'GCN', 'THƯ MỜI': 'TM', 'CHUYỂN ĐIỂM': 'CĐ'}
DOC_TYPE_CODES={**DOCUMENT_TYPE_CODES,'CÔNG VĂN':'CV','CÔNG VĂN ĐẾN':'CVDEN','CÔNG VĂN ĐI':'CV','HỒ SƠ KHÁC':'HS'}
OUTGOING_TYPES={'CÔNG VĂN',*DOCUMENT_TYPE_CODES}
ARCHIVE_FOLDERS={'CÔNG VĂN ĐẾN','CÔNG VĂN ĐI','HỒ SƠ KHÁC',*DOCUMENT_TYPE_CODES}
SCHOOL_DEPARTMENTS=[
 ('KQTKD_MKT','Khoa Quản trị kinh doanh – Marketing'),
 ('KTCNHKT','Khoa Tài chính – Ngân hàng – Kế toán'),
 ('KLUAT','Khoa Luật'),
 ('KDL_NH_KS','Khoa Du lịch – Nhà hàng – Khách sạn'),
 ('KKHSK','Khoa Khoa học sức khoẻ'),
 ('KKTCN','Khoa Kỹ thuật Công nghệ'),
 ('KNN','Khoa Ngôn ngữ'),
 ('KKHLN','Khoa Khoa học liên ngành'),
 ('VCNTT_AI','Viện Công nghệ tiên tiến và Trí tuệ nhân tạo'),
 ('VVHDN','Viện Văn hoá Doanh nghiệp'),
 ('VDTSĐH','Viện Đào tạo Sau đại học'),
 ('VLKGD_DTTX','Viện Liên kết Giáo dục và Đào tạo từ xa'),
 ('TTCN','Trung tâm Công nghệ'),
 ('TTTT_PTTS','Trung tâm Truyền thông và Phát triển Tuyển sinh'),
 ('TTHL','Trung tâm Học liệu'),
 ('VP','Văn phòng Trường'),
 ('PDT','Phòng Đào tạo'),
 ('TCKT','Phòng Tài chính – Kế toán'),
 ('PQLKH','Phòng Quản lý Khoa học'),
 ('PCTSV_XH','Phòng Công tác sinh viên và Xã hội'),
 ('PKT_QLCL','Phòng Khảo thí và Quản lý chất lượng'),
 ('PHT_PT','Phòng Hợp tác và Phát triển'),
 ('TCKH','Tạp chí Khoa học Trường Đại học Hùng Vương Thành phố Hồ Chí Minh'),
 ('VUOM_DHV','Vườn ươm khởi nghiệp DHV'),
 ('CBT','Chi bộ Trường'),
 ('VDMST_KN','Viện Đổi mới sáng tạo và Khởi nghiệp'),
]
INTERNAL_UNITS={name for _,name in SCHOOL_DEPARTMENTS}
def normalize_search(value:str)->str:
 value=(value or '').replace('Đ','D').replace('đ','d').casefold()
 return ' '.join(''.join(c for c in unicodedata.normalize('NFD',value) if unicodedata.category(c)!='Mn').split())
def number_symbol(doc_type:str,number:int,when:date,direction:Optional[str]=None):
 if direction=='IN' or doc_type=='CÔNG VĂN ĐẾN':return f'{number:02d}/{when.year}/CVDEN'
 if doc_type in {'CÔNG VĂN','CÔNG VĂN ĐI','CÔNG VĂN NỘI BỘ'}:return f'{number:02d}/{when.year}/ĐHHV'
 code=DOC_TYPE_CODES.get(doc_type,doc_type)
 suffix=code if code.endswith('-HĐT') else code+'-ĐHHV'
 return f'{number:02d}/{when.year}/{suffix}'
def system_number_floor(s:Session,year:int,category:str)->int:
 # Archive folders and workflow document types share one numbering category.
 category_sql="CASE WHEN direction='IN' OR doc_type='CÔNG VĂN ĐẾN' THEN 'CÔNG VĂN ĐẾN' WHEN doc_type IN ('CÔNG VĂN','CÔNG VĂN ĐI','CÔNG VĂN NỘI BỘ') THEN 'CÔNG VĂN ĐI' ELSE doc_type END"
 highest=s.connection().exec_driver_sql(f"SELECT COALESCE(MAX(number),0) FROM managed_documents WHERE {category_sql}=? AND COALESCE(archive_year,CAST(strftime('%Y',COALESCE(issued_date,received_date,created_at)) AS INTEGER))=?",(category,year)).scalar_one()
 if category=='CÔNG VĂN ĐẾN':
  highest=max(highest,s.connection().exec_driver_sql("SELECT COALESCE(MAX(number),0) FROM incoming_numbers WHERE year=?",(year,)).scalar_one())
 current=s.connection().exec_driver_sql("SELECT COALESCE(MAX(current),0) FROM sequences WHERE year=? AND doc_type=?",(year,category)).scalar_one()
 return max(highest,current)

def allocate_archive_number(s:Session,year:int,doc_type:str)->int:
 category='CÔNG VĂN ĐI' if doc_type in {'CÔNG VĂN','CÔNG VĂN NỘI BỘ'} else doc_type
 floor=system_number_floor(s,year,category)
 sql='''INSERT INTO sequences (year,doc_type,current,prefix) VALUES (?,?,?,?)
 ON CONFLICT(year,doc_type) DO UPDATE SET current=MAX(sequences.current,excluded.current-1)+1
 RETURNING current'''
 return int(s.connection().exec_driver_sql(sql,(year,category,floor+1,DOC_TYPE_CODES.get(category,'HS'))).scalar_one())
def file_title_part(name:str)->str:
 stem=Path(name or '').stem.strip()
 return re.sub(r'^[A-ZĐ]+\s+\d+-\d{2},\s*\d{2}-\d{2}\s*','',stem,flags=re.IGNORECASE).strip()
def safe_file_title(title:str)->str:
 value=re.sub(r'[\\/:*?"<>|\0]+',' ',(title or '').strip())
 value=' '.join(value.split()).strip(' .')
 return (value or 'Văn bản')[:180].rstrip(' .')
def numbered_file_name(symbol:str,title:str)->str:
 return f'{safe_file_title(symbol)} {safe_file_title(title)}.pdf'
def document_file_name(title:str,symbol:Optional[str]=None)->str:
 return numbered_file_name(symbol,title) if symbol else f'{safe_file_title(title)}.pdf'
def clean_workflow_suffixes(name:str)->str:
 if not name:return name
 suffix=Path(name).suffix or '.pdf';stem=Path(name).stem
 while True:
  cleaned=re.sub(r'-(?:office-opinion-signed|office-opinion|signed|sealed)$','',stem,flags=re.IGNORECASE)
  if cleaned==stem:break
  stem=cleaned
 return f'{stem}{suffix}'
def resolve_managed_file(d)->Optional[Path]:
 path=Path(d.file_path) if d.file_path else None
 if path and path.is_file():return path
 candidates=[]
 if getattr(d,'id',None):candidates.extend(ARCHIVE_ROOT.rglob(f'{d.id:06d}-*'))
 if d.file_name:
  candidates.extend(ARCHIVE_ROOT.rglob(d.file_name));candidates.extend(STORE.rglob(Path(d.file_name).name))
 return next((candidate for candidate in candidates if candidate.is_file()),None)
def resolve_source_file(source:DocumentSource,d:Document)->Optional[Path]:
 path=Path(source.file_path) if source.file_path else None
 if path and path.is_file():return path
 candidates=[]
 if getattr(d,'id',None):candidates.extend(ARCHIVE_ROOT.rglob(f'{d.id:06d}-*'))
 for name in (source.file_name,d.file_name):
  if name:
   candidates.extend(ARCHIVE_ROOT.rglob(Path(name).name))
   candidates.extend(STORE.rglob(Path(name).name))
 return next((candidate for candidate in candidates if candidate.is_file()),None)
engine=create_engine(f"sqlite:///{ROOT/'dhv.db'}",connect_args={'check_same_thread':False});DB=sessionmaker(bind=engine,expire_on_commit=False)
SECRET=os.getenv('SECRET_KEY','change-me-production');pwd=CryptContext(schemes=['pbkdf2_sha256'],deprecated='auto');oauth=OAuth2PasswordBearer(tokenUrl='/api/auth/login')
class Base(DeclarativeBase):pass
class Department(Base):
 __tablename__='departments';id:Mapped[int]=mapped_column(primary_key=True);name:Mapped[str]=mapped_column(unique=True);code:Mapped[str]=mapped_column(unique=True);email:Mapped[str]=mapped_column(default='')
class User(Base):
 __tablename__='users';id:Mapped[int]=mapped_column(primary_key=True);username:Mapped[str]=mapped_column(unique=True);full_name:Mapped[str];password_hash:Mapped[str];role:Mapped[str];department_id:Mapped[Optional[int]]=mapped_column(ForeignKey('departments.id'));active:Mapped[bool]=mapped_column(Boolean,default=True)
class Sequence(Base):
 __tablename__='sequences';__table_args__=(UniqueConstraint('year','doc_type'),);id:Mapped[int]=mapped_column(primary_key=True);year:Mapped[int];doc_type:Mapped[str];current:Mapped[int]=mapped_column(default=0);prefix:Mapped[str]
class Document(Base):
 __tablename__='managed_documents';id:Mapped[int]=mapped_column(primary_key=True);direction:Mapped[str];doc_type:Mapped[str];number:Mapped[Optional[int]];symbol:Mapped[Optional[str]];title:Mapped[str];summary:Mapped[str]=mapped_column(Text,default='');issuing_agency:Mapped[str]=mapped_column(default='');issued_date:Mapped[Optional[date]]=mapped_column(Date);received_date:Mapped[Optional[date]]=mapped_column(Date);archive_year:Mapped[Optional[int]];status:Mapped[str]=mapped_column(default='DRAFT');revision:Mapped[int]=mapped_column(default=1);priority:Mapped[str]=mapped_column(default='Bình thường');owner_id:Mapped[int]=mapped_column(ForeignKey('users.id'));department_id:Mapped[Optional[int]]=mapped_column(ForeignKey('departments.id'));assignee_ids:Mapped[str]=mapped_column(default='[]');file_name:Mapped[Optional[str]];file_path:Mapped[Optional[str]];ocr_text:Mapped[str]=mapped_column(Text,default='');search_text:Mapped[str]=mapped_column(Text,default='');reserved:Mapped[bool]=mapped_column(Boolean,default=False);signed_personal:Mapped[bool]=mapped_column(Boolean,default=False);sealed:Mapped[bool]=mapped_column(Boolean,default=False);signature_hash:Mapped[Optional[str]];created_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow);updated_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow,onupdate=datetime.utcnow)
class Action(Base):
 __tablename__='actions';id:Mapped[int]=mapped_column(primary_key=True);document_id:Mapped[int]=mapped_column(ForeignKey('managed_documents.id'));user_id:Mapped[int]=mapped_column(ForeignKey('users.id'));action:Mapped[str];comment:Mapped[str]=mapped_column(default='');created_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow)
class DigitalSignature(Base):
 __tablename__='digital_signatures';id:Mapped[int]=mapped_column(primary_key=True);document_id:Mapped[int]=mapped_column(ForeignKey('managed_documents.id'));user_id:Mapped[int]=mapped_column(ForeignKey('users.id'));page:Mapped[int]=mapped_column(default=1);x_percent:Mapped[float]=mapped_column(Float);y_percent:Mapped[float]=mapped_column(Float);file_hash:Mapped[str];created_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow)
class DigitalSeal(Base):
 __tablename__='digital_seals';id:Mapped[int]=mapped_column(primary_key=True);document_id:Mapped[int]=mapped_column(ForeignKey('managed_documents.id'));user_id:Mapped[int]=mapped_column(ForeignKey('users.id'));seal_type:Mapped[str];page:Mapped[int]=mapped_column(default=1);x_percent:Mapped[float]=mapped_column(Float);y_percent:Mapped[float]=mapped_column(Float);file_hash:Mapped[str];created_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow)
class DocumentSource(Base):
 __tablename__='document_sources';id:Mapped[int]=mapped_column(primary_key=True);document_id:Mapped[int]=mapped_column(ForeignKey('managed_documents.id'),unique=True);file_path:Mapped[str];file_name:Mapped[str]
class IncomingNumber(Base):
 __tablename__='incoming_numbers';__table_args__=(UniqueConstraint('year','doc_type','number'),);id:Mapped[int]=mapped_column(primary_key=True);document_id:Mapped[int]=mapped_column(ForeignKey('managed_documents.id'),unique=True);year:Mapped[int];doc_type:Mapped[str];number:Mapped[int]
class DeletedRecord(Base):
 __tablename__='deleted_records';id:Mapped[int]=mapped_column(primary_key=True);document_id:Mapped[int]=mapped_column(ForeignKey('managed_documents.id'),unique=True);previous_status:Mapped[str];deleted_by:Mapped[int]=mapped_column(ForeignKey('users.id'));reason:Mapped[str];deleted_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow)
class EmailSetting(Base):
 __tablename__='email_settings';id:Mapped[int]=mapped_column(primary_key=True);smtp_host:Mapped[str]=mapped_column(default='');smtp_port:Mapped[int]=mapped_column(default=587);username:Mapped[str]=mapped_column(default='');sender_name:Mapped[str]=mapped_column(default='Đại học Hùng Vương');sender_email:Mapped[str]=mapped_column(default='');signature_html:Mapped[str]=mapped_column(Text,default='');reminder_subject:Mapped[str]=mapped_column(default='Nhắc hạn nộp báo cáo: {title}');reminder_html:Mapped[str]=mapped_column(Text,default='<p>Đơn vị chưa nộp báo cáo cho hồ sơ <b>{title}</b>.</p><p>Hạn nộp: <b>{deadline}</b>.</p>');use_tls:Mapped[bool]=mapped_column(Boolean,default=True);enabled:Mapped[bool]=mapped_column(Boolean,default=False);updated_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow,onupdate=datetime.utcnow)
class RolePermission(Base):
 __tablename__='role_permissions';__table_args__=(UniqueConstraint('role','permission'),);id:Mapped[int]=mapped_column(primary_key=True);role:Mapped[str];permission:Mapped[str];enabled:Mapped[bool]=mapped_column(Boolean,default=True)
class IncomingReview(Base):
 __tablename__='incoming_reviews';__table_args__=(UniqueConstraint('document_id','revision'),);id:Mapped[int]=mapped_column(primary_key=True);document_id:Mapped[int]=mapped_column(ForeignKey('managed_documents.id'));revision:Mapped[int];note:Mapped[str]=mapped_column(Text);assignments:Mapped[str]=mapped_column(Text);decisions:Mapped[str]=mapped_column(Text,default='{}');pdf_path:Mapped[str];created_by:Mapped[int]=mapped_column(ForeignKey('users.id'));created_at:Mapped[datetime]=mapped_column(DateTime,default=datetime.utcnow)
Base.metadata.create_all(engine)
with engine.begin() as conn:
 if 'email' not in {row[1] for row in conn.exec_driver_sql('PRAGMA table_info(departments)')}:conn.exec_driver_sql("ALTER TABLE departments ADD COLUMN email TEXT NOT NULL DEFAULT ''")
 columns={row[1] for row in conn.exec_driver_sql('PRAGMA table_info(managed_documents)')}
 if 'search_text' not in columns:conn.exec_driver_sql("ALTER TABLE managed_documents ADD COLUMN search_text TEXT NOT NULL DEFAULT ''")
 if 'revision' not in columns:conn.exec_driver_sql("ALTER TABLE managed_documents ADD COLUMN revision INTEGER NOT NULL DEFAULT 1")
 if 'archive_year' not in columns:conn.exec_driver_sql("ALTER TABLE managed_documents ADD COLUMN archive_year INTEGER")
 conn.exec_driver_sql("UPDATE managed_documents SET status='PENDING_OUT_BGH_APPROVAL' WHERE status IN ('PENDING_CLERK_REVIEW','PENDING_OUT_BGH')")
 conn.exec_driver_sql("UPDATE managed_documents SET status='READY_FOR_CLERK' WHERE status='OUT_BGH_SIGNED'")
 email_columns={row[1] for row in conn.exec_driver_sql('PRAGMA table_info(email_settings)')}
 if 'signature_html' not in email_columns:conn.exec_driver_sql("ALTER TABLE email_settings ADD COLUMN signature_html TEXT NOT NULL DEFAULT ''")
 if 'reminder_subject' not in email_columns:conn.exec_driver_sql("ALTER TABLE email_settings ADD COLUMN reminder_subject TEXT NOT NULL DEFAULT 'Nhắc hạn nộp báo cáo: {title}'")
 if 'reminder_html' not in email_columns:conn.exec_driver_sql("ALTER TABLE email_settings ADD COLUMN reminder_html TEXT NOT NULL DEFAULT '<p>Đơn vị chưa nộp báo cáo cho hồ sơ <b>{title}</b>.</p><p>Hạn nộp: <b>{deadline}</b>.</p>'")
def db():
 s=DB()
 try:yield s
 finally:s.close()
def seed():
 s=DB()
 if not s.query(User).count():
  ds=[Department(name='Ban Giám hiệu',code='BGH')]+[Department(name=name,code=code) for code,name in SCHOOL_DEPARTMENTS];s.add_all(ds);s.flush();by_code={d.code:d for d in ds}
  for u,n,r,c in [('admin','Quản trị hệ thống','ADMIN','BGH'),('hieu_truong','Hiệu trưởng Nguyễn Văn A','BGH','BGH'),('truong_vp','Chánh Văn phòng','OFFICE_HEAD','VP'),('van_thu','Chuyên viên Văn thư','CLERK','VP'),('phong_dao_tao','Chuyên viên Đào tạo','DEPARTMENT','PDT')]:s.add(User(username=u,full_name=n,role=r,department_id=by_code[c].id,password_hash=pwd.hash('123456')))
  for t,p in [('QUYẾT ĐỊNH','QĐ-DHV'),('THÔNG BÁO','TB-DHV'),('CÔNG VĂN','DHV'),('KẾ HOẠCH','KH-DHV')]:s.add(Sequence(year=date.today().year,doc_type=t,prefix=p))
 s.commit()
 s.close()
def sync_school_departments():
 """Keep one canonical unit directory while preserving existing ids and links."""
 aliases={'VP':{'Văn phòng'},'TCKT':{'Phòng Tài chính - Kế toán'},'PCTSV_XH':{'Phòng Công tác sinh viên'},'PQLKH':{'Phòng Quản ký Khoa học'}}
 with DB() as s:
  existing=s.query(Department).all();by_code={d.code:d for d in existing};by_name={normalize_search(d.name):d for d in existing}
  for code,name in SCHOOL_DEPARTMENTS:
   dep=by_code.get(code) or by_name.get(normalize_search(name))
   if not dep:
    dep=next((by_name.get(normalize_search(alias)) for alias in aliases.get(code,set()) if by_name.get(normalize_search(alias))),None)
   if dep:dep.code=code;dep.name=name
   else:dep=Department(code=code,name=name);s.add(dep)
   by_code[code]=dep;by_name[normalize_search(name)]=dep
  s.flush()
  # Merge former names that have a clear successor in the canonical directory.
  legacy_targets={'Phòng Truyền thông':'TTTT_PTTS','Trung tâm Tuyển sinh':'TTTT_PTTS'}
  for old in list(s.query(Department).all()):
   target_code=legacy_targets.get(old.name);target=by_code.get(target_code) if target_code else None
   if not target or old.id==target.id:continue
   if not target.email and old.email:target.email=old.email
   s.query(User).filter(User.department_id==old.id).update({User.department_id:target.id},synchronize_session=False)
   s.query(Document).filter(Document.department_id==old.id).update({Document.department_id:target.id},synchronize_session=False)
   for document in s.query(Document).filter(Document.assignee_ids.like(f'%{old.id}%')):
    try:document.assignee_ids=json.dumps(sorted({target.id if value==old.id else value for value in json.loads(document.assignee_ids)}))
    except (TypeError,ValueError):pass
   for review in s.query(IncomingReview).filter(IncomingReview.assignments.like(f'%{old.id}%')):
    assignments=json.loads(review.assignments);changed=False
    for assignment in assignments:
     if assignment.get('department_id')==old.id:assignment.update(department_id=target.id,name=target.name,email=target.email);changed=True
    if changed:
     unique={item['department_id']:item for item in assignments};review.assignments=json.dumps(list(unique.values()),ensure_ascii=False)
     decisions=json.loads(review.decisions);old_decision=decisions.pop(str(old.id),None)
     if old_decision and str(target.id) not in decisions:decisions[str(target.id)]=old_decision
     review.decisions=json.dumps(decisions,ensure_ascii=False)
   s.delete(old)
  s.commit()
seed();sync_school_departments();app=FastAPI(title='DHV - Quản lý văn bản',version='1.0');app.add_middleware(CORSMiddleware,allow_origins=['*'],allow_methods=['*'],allow_headers=['*'])
with DB() as admin_session:
 admin_account=admin_session.query(User).filter_by(username='admin').first()
 if admin_account and (not admin_account.active or admin_account.role!='ADMIN'):
  admin_account.active=True;admin_account.role='ADMIN';admin_session.commit()
class Login(BaseModel):username:str;password:str
class AIChatMessage(BaseModel):role:str;content:str
class AIChatRequest(BaseModel):message:str;history:list[AIChatMessage]=[]
class AIFileSearchRequest(BaseModel):query:str
class AIAnalysisRequest(BaseModel):message:str;document_ids:list[int]=Field(min_length=1,max_length=5);history:list[AIChatMessage]=[]
class AIReportRequest(BaseModel):
 requirement:str=Field(min_length=5,max_length=4000)
 document_ids:list[int]=Field(min_length=1,max_length=15)
 audience:str='Ban Giám hiệu'
 detail_level:str='FULL'
 outline:list[str]=Field(default_factory=list,max_length=12)
 instruction:str=''
 history:list[AIChatMessage]=[]
class AIReportExport(BaseModel):title:str='Báo cáo tổng hợp công văn';content:str
class Act(BaseModel):action:str;comment:str='';department_ids:list[int]=[]
class Reserve(BaseModel):doc_type:str;year:int=date.today().year;quantity:int=1;department_id:Optional[int]=None
class RenameFile(BaseModel):name:str
class UserConfig(BaseModel):full_name:str;role:str;department_id:Optional[int]=None;active:bool=True
class EmailConfig(BaseModel):smtp_host:str='';smtp_port:int=587;username:str='';sender_name:str='Đại học Hùng Vương';sender_email:str='';signature_html:str='';reminder_subject:str='Nhắc hạn nộp báo cáo: {title}';reminder_html:str='<p>Đơn vị chưa nộp báo cáo cho hồ sơ <b>{title}</b>.</p><p>Hạn nộp: <b>{deadline}</b>.</p>';use_tls:bool=True;enabled:bool=False
class ArchiveReviewDecision(BaseModel):decision:Literal['APPROVE','RETURN'];note:str
class PermissionConfig(BaseModel):role:str;permissions:list[str]
def current(token=Depends(oauth),s:Session=Depends(db)):
 try:uid=int(jwt.decode(token,SECRET,algorithms=['HS256'])['sub'])
 except (JWTError,KeyError):raise HTTPException(401,'Phiên đăng nhập không hợp lệ')
 u=s.get(User,uid)
 if not u or not u.active:raise HTTPException(401,'Tài khoản bị khóa')
 return u
PERMISSION_LABELS={'VIEW_ALL':'Xem toàn bộ hồ sơ','CREATE':'Tạo và tiếp nhận văn bản','EDIT':'Sửa và bổ sung văn bản','RESERVE_NUMBER':'Xin số trước','ASSIGN_NUMBER':'Cấp số văn bản','OFFICE_OPINION':'Ý kiến Chánh Văn phòng','INTERNAL_APPROVE':'Phê duyệt văn bản nội bộ','SEND_BGH':'Trình Ban Giám hiệu','BGH_DECIDE':'Duyệt hoặc từ chối','FORWARD':'Chuyển đơn vị xử lý','SIGN':'Ký số','SEAL':'Đóng dấu','EMAIL':'Phát hành qua email','ARCHIVE':'Lưu trữ và OCR','RENAME':'Đổi tên file','DELETE':'Xóa có lý do','RESTORE':'Khôi phục văn bản','MANAGE_USERS':'Quản lý người dùng','MANAGE_PERMISSIONS':'Cấu hình phân quyền','MANAGE_EMAIL':'Cấu hình email','VIEW_ACTIVITY':'Xem activity log','AI_CHAT':'Dùng trợ lý AI tra cứu'}
ROLE_DEFAULTS={'ADMIN':set(PERMISSION_LABELS),'BGH':{'BGH_DECIDE','SIGN'},'OFFICE_HEAD':{'VIEW_ALL','OFFICE_OPINION','SEND_BGH','INTERNAL_APPROVE','SIGN','AI_CHAT'},'CLERK':{'VIEW_ALL','CREATE','EDIT','ASSIGN_NUMBER','FORWARD','SEAL','EMAIL','ARCHIVE','RENAME','DELETE','AI_CHAT'},'DEPARTMENT':{'CREATE','EDIT','SIGN','RESERVE_NUMBER'},'DEPARTMENT_HEAD':{'CREATE','EDIT','SIGN','RESERVE_NUMBER'}}
def migrate_legacy_department_permissions():
 """Replace the old read/AI-only unit profile without overriding later admin choices."""
 with DB() as session:
  for role in ('DEPARTMENT','DEPARTMENT_HEAD'):
   rows={row.permission:row for row in session.query(RolePermission).filter_by(role=role)}
   legacy=rows.get('AI_CHAT') and rows['AI_CHAT'].enabled and all(not rows.get(code) or not rows[code].enabled for code in ('CREATE','EDIT','SIGN'))
   if not legacy:continue
   for code in ('CREATE','EDIT','SIGN'):
    row=rows.get(code)
    if not row:row=RolePermission(role=role,permission=code);session.add(row)
    row.enabled=True
   rows['AI_CHAT'].enabled=False
  session.commit()
migrate_legacy_department_permissions()
def has_permission(s,u,permission):
 if permission=='RESERVE_NUMBER':return u.role in {'ADMIN','CLERK','OFFICE_HEAD'} or (u.role in {'DEPARTMENT','DEPARTMENT_HEAD'} and u.department_id is not None)
 if u.role=='ADMIN':return True
 if permission not in ROLE_DEFAULTS.get(u.role,set()):return False
 row=s.query(RolePermission).filter_by(role=u.role,permission=permission).first()
 return row.enabled if row else permission in ROLE_DEFAULTS.get(u.role,set())
def require_permission(s,u,permission,message='Bạn không có quyền thực hiện thao tác này'):
 if not has_permission(s,u,permission):raise HTTPException(403,message)
def require_office_module(u):
 if u.role not in {'ADMIN','CLERK','OFFICE_HEAD'}:raise HTTPException(403,'Module này chỉ dành cho Văn phòng Trường')

def view(u,d,s=None):
 if u.role=='BGH':
  if d.owner_id==u.id:return True
  if d.direction=='OUT' and d.status in {'PENDING_OUT_BGH_APPROVAL','PENDING_OUT_BGH_SIGN'}:return True
  if s is None:return False
  if d.direction=='IN' and s.query(IncomingReview.id).filter_by(document_id=d.id).first():return True
  return s.query(Action.id).filter_by(document_id=d.id,user_id=u.id).first() is not None
 return (has_permission(s,u,'VIEW_ALL') if s else u.role in {'ADMIN','OFFICE_HEAD','CLERK'}) or u.id==d.owner_id or (u.department_id is not None and (u.department_id==d.department_id or u.department_id in json.loads(d.assignee_ids or '[]')))
def require_document_access(s,u,d):
 if not view(u,d,s):raise HTTPException(404,'Không tìm thấy văn bản')
def out(d):return {c.name:getattr(d,c.name) for c in d.__table__.columns if c.name!='search_text'}
def refresh_search(d):d.search_text=normalize_search(' '.join(str(x or '') for x in (d.symbol,d.title,d.summary,d.issuing_agency,d.file_name,d.ocr_text)))
class EmailHTMLSanitizer(HTMLParser):
 allowed={'p','div','br','b','strong','i','em','u','ul','ol','li','blockquote','span','font','img'}
 styles={'font-weight','font-style','text-decoration','font-size','font-family','color','margin-left','text-align'}
 def __init__(self):super().__init__(convert_charrefs=True);self.parts=[]
 def handle_starttag(self,tag,attrs):
  if tag not in self.allowed:return
  clean=[]
  for key,value in attrs:
   if tag=='img' and key=='src' and (value.startswith('data:image/png;base64,') or value.startswith('data:image/jpeg;base64,') or value.startswith('data:image/gif;base64,')):clean.append(('src',value))
   elif tag=='img' and key in {'alt','width','height'} and re.fullmatch(r'[\w .,%()\-À-ỹ]*',value or ''):clean.append((key,value))
   elif tag=='font' and key in {'size','color','face'} and re.fullmatch(r'[\w #(),.%\-À-ỹ]*',value or ''):clean.append((key,value))
   elif key=='style':
    declarations=[]
    for item in (value or '').split(';'):
     name,sep,val=item.partition(':');name=name.strip().lower();val=val.strip()
     if sep and name in self.styles and not re.search(r'url|expression|javascript',val,re.I):declarations.append(f'{name}:{val}')
    if declarations:clean.append(('style',';'.join(declarations)))
  rendered=''.join(f' {k}="{html.escape(v,quote=True)}"' for k,v in clean);self.parts.append(f'<{tag}{rendered}>')
 def handle_endtag(self,tag):
  if tag in self.allowed and tag not in {'br','img'}:self.parts.append(f'</{tag}>')
 def handle_data(self,data):self.parts.append(html.escape(data))

def clean_email_html(value:str)->str:
 parser=EmailHTMLSanitizer();parser.feed(value or '');return ''.join(parser.parts).strip()
def constrain_signature_images(value:str,max_width:int=100)->str:
 def resize(match):
  tag=match.group(0);source=re.search(r'src="data:image/(?:png|jpeg|gif);base64,([A-Za-z0-9+/=]+)"',tag,re.I)
  if not source:return tag
  try:
   with Image.open(io.BytesIO(base64.b64decode(source.group(1),validate=True))) as picture:natural_width,natural_height=picture.size
  except Exception:return tag
  width=min(natural_width,max_width);height=max(1,round(width*natural_height/natural_width))
  tag=re.sub(r'\s(?:width|height|style)="[^"]*"','',tag,flags=re.I)
  return tag[:-1]+f' width="{width}" height="{height}" style="width:{width}px;height:auto;max-width:100%">'
 return re.sub(r'<img\b[^>]*>',resize,value or '',flags=re.I)
def has_email_content(value:str)->bool:return bool(re.sub(r'\s+',' ',re.sub(r'<[^>]+>',' ',value or '')).strip() or re.search(r'<img\b',value or '',re.I))
def email_plain_text(value:str)->str:
 value=re.sub(r'<(?:br|/p|/div|/li)>', '\n',value or '',flags=re.I);return html.unescape(re.sub(r'<[^>]+>','',value)).strip()
def prepare_inline_images(value:str):
 images=[]
 def replace(match):
  subtype={'jpeg':'jpeg','png':'png','gif':'gif'}[match.group(1)];cid=f'email-image-{uuid.uuid4().hex}'
  try:data=base64.b64decode(match.group(2),validate=True)
  except ValueError:return ''
  if len(data)>1024*1024:return ''
  images.append((cid,subtype,data));return f'cid:{cid}'
 return re.sub(r'data:image/(png|jpeg|gif);base64,([A-Za-z0-9+/=]+)',replace,value,flags=re.I),images

def gmail_draft(d,recipients,subject,content,signature=''):
 if not recipients or any(not re.fullmatch(r'[^\s@,;<>]+@[^\s@,;<>]+\.[^\s@,;<>]+',address) for address in recipients):raise HTTPException(422,'Nhập email người nhận hợp lệ')
 if not subject.strip() or not has_email_content(content):raise HTTPException(422,'Tiêu đề và nội dung email là bắt buộc')
 if not d.file_path or not Path(d.file_path).is_file():raise HTTPException(400,'Văn bản chưa có tệp để đính kèm')
 body=email_plain_text(content)+'\n\n'+email_plain_text(signature)
 url='https://mail.google.com/mail/?'+urlencode({'view':'cm','fs':'1','authuser':'ngcphnglinhp6.7.2000@gmail.com','to':','.join(recipients),'su':subject.strip(),'body':body.strip()})
 return {'email_status':'draft','gmail_url':url,'gmail_account':'ngcphnglinhp6.7.2000@gmail.com','attachment_url':f'/documents/{d.id}/file','attachment_name':d.file_name,'document':out(d)}

def send_document_email(s:Session,d:Document,recipients:list[str],subject:str,content_html:str,signature_html:str):
 setting=s.query(EmailSetting).first()
 if not setting or not setting.enabled:raise HTTPException(409,'Gửi email đang tắt trong phần cấu hình hệ thống')
 password=os.getenv('SMTP_PASSWORD','').replace(' ','')
 if not password:raise HTTPException(503,'Máy chủ chưa được cấu hình biến môi trường SMTP_PASSWORD')
 if not d.file_path or not Path(d.file_path).is_file():raise HTTPException(400,'Văn bản chưa có tệp PDF để gửi')
 subject=subject.strip();content_html=clean_email_html(content_html);signature_html=constrain_signature_images(clean_email_html(signature_html))
 if not subject:raise HTTPException(422,'Tiêu đề email là bắt buộc')
 if not has_email_content(content_html):raise HTTPException(422,'Nội dung email là bắt buộc')
 if not has_email_content(signature_html):raise HTTPException(422,'Chữ ký email là bắt buộc')
 full_html=f'<div>{content_html}</div><div style="margin-top:24px">{signature_html}</div>';full_html,inline_images=prepare_inline_images(full_html)
 message=EmailMessage();message['Subject']=subject;message['From']=f'{setting.sender_name} <{setting.sender_email}>';message['To']=', '.join(recipients)
 message.set_content(f'{email_plain_text(content_html)}\n\n{email_plain_text(signature_html)}');message.add_alternative(full_html,subtype='html')
 html_part=message.get_payload()[-1]
 for cid,subtype,data in inline_images:html_part.add_related(data,maintype='image',subtype=subtype,cid=f'<{cid}>',filename=f'{cid}.{subtype}')
 file_path=Path(d.file_path);message.add_attachment(file_path.read_bytes(),maintype='application',subtype='pdf',filename=d.file_name or file_path.name)
 try:
  with smtplib.SMTP(setting.smtp_host,setting.smtp_port,timeout=120) as smtp:
   smtp.ehlo()
   if setting.use_tls:smtp.starttls();smtp.ehlo()
   smtp.login(setting.username,password);smtp.send_message(message)
 except smtplib.SMTPAuthenticationError as exc:
  raise HTTPException(502,'Gmail từ chối đăng nhập. Hãy kiểm tra email gửi và tạo lại App Password') from exc
 except smtplib.SMTPRecipientsRefused as exc:
  rejected=', '.join(exc.recipients.keys())
  raise HTTPException(422,f'Gmail từ chối địa chỉ người nhận: {rejected}') from exc
 except smtplib.SMTPSenderRefused as exc:
  raise HTTPException(422,'Gmail từ chối địa chỉ người gửi. Email người gửi phải trùng tài khoản SMTP') from exc
 except smtplib.SMTPDataError as exc:
  raise HTTPException(502,f'Gmail từ chối nội dung thư (mã SMTP {exc.smtp_code})') from exc
 except (OSError,smtplib.SMTPException) as exc:
  raise HTTPException(502,f'Không thể gửi email qua SMTP ({type(exc).__name__})') from exc
with DB() as search_session:
 for search_document in search_session.query(Document):
  old_name=search_document.file_name or '';clean_name=clean_workflow_suffixes(old_name)
  if old_name:clean_name=document_file_name(search_document.title,search_document.symbol)
  if clean_name!=old_name:
   path=Path(search_document.file_path) if search_document.file_path else None
   if path and path.exists() and path.name.endswith(old_name):
    target=path.with_name(path.name[:-len(old_name)]+clean_name)
    if not target.exists():path.rename(target);search_document.file_path=str(target)
   search_document.file_name=clean_name
  refresh_search(search_document)
 search_session.commit()
def log(s,d,u,a,c=''):s.add(Action(document_id=d.id,user_id=u.id,action=a,comment=c))
def pdf_point(page,x_percent,y_percent):
 return page.rect.width*x_percent/100,page.rect.height*y_percent/100
def render_signature(source:Path,x_percent:int,y_percent:int,page_number:int)->Path:
 doc=pymupdf.open(source);page=doc[max(0,min(len(doc)-1,page_number-1))];cx,cy=pdf_point(page,x_percent,y_percent);w,h=150,62;visual=pymupdf.Rect(cx-w/2,cy-h/2,cx+w/2,cy+h/2);box=visual*page.derotation_matrix;page.draw_oval(box,color=(.75,.04,.08),width=2.2,overlay=True);page.insert_textbox(box,'\nDA KY SO - GIA LAP\nDH HUNG VUONG',fontname='helv',fontsize=8,color=(.75,.04,.08),align=1,rotate=page.rotation,overlay=True);target=STORE/f'{uuid.uuid4()}-signed.pdf';doc.save(target,garbage=4,deflate=True);doc.close();return target
class NumberRegion(BaseModel):
 x:float=Field(ge=0,lt=100)
 y:float=Field(ge=0,lt=100)
 width:float=Field(gt=0,le=100)
 height:float=Field(gt=0,le=100)
class NumberPlacement(BaseModel):
 number:NumberRegion
 date:NumberRegion

def _number_placement_from_words(words:list[tuple],page_width:float,page_height:float)->Optional[NumberPlacement]:
 # Some PDFs split visually adjacent words into unrelated text blocks. Group by
 # their physical baseline so the detector still sees one printed line.
 lines=[]
 for word in sorted((word for word in words if len(word)>=5 and str(word[4]).strip()),key=lambda item:((item[1]+item[3])/2,item[0])):
  center=(word[1]+word[3])/2
  line=next((candidate for candidate in lines if abs(candidate[0]-center)<=max(4,(word[3]-word[1])*.45)),None)
  if line:line[1].append(word);line[0]=(line[0]*(len(line[1])-1)+center)/len(line[1])
  else:lines.append([center,[word]])
 number_rect=date_rect=None
 for _,line in lines:
  line.sort(key=lambda word:word[0])
  normalized=[normalize_search(str(word[4])).strip('.,:;()') for word in line]
  if number_rect is None:
   for index,value in enumerate(normalized):
    raw_label=str(line[index][4]).strip()
    if value!='so' and not (raw_label.casefold().startswith('s') and ':' in raw_label):continue
    label=line[index];candidate=line[index+1] if index+1<len(line) else None
    if candidate:
     raw=str(candidate[4]).strip();match=re.match(r'(\d+)',raw)
     if match:
      ratio=max(.12,min(1,len(match.group(1))/max(1,len(raw))))
      x0,x1=candidate[0]-1,candidate[0]+(candidate[2]-candidate[0])*ratio+2
     elif raw.startswith('/'):
      x1=candidate[0]-1;x0=max(label[2]+2,x1-max(28,label[3]-label[1])*2.5)
     else:x0=label[2]+3;x1=min(x0+max(32,(label[3]-label[1])*2.8),candidate[0]-1)
    else:x0=label[2]+3;x1=x0+max(32,(label[3]-label[1])*2.8)
    y0=label[1]-2;y1=label[3]+2
    if x1>x0+3:number_rect=(x0,y0,x1,y1)
    break
  if date_rect is None:
   for index,value in enumerate(normalized):
    if value!='ngay':continue
    end=index
    for cursor in range(index+1,len(line)):
     end=cursor
     if normalized[cursor]=='nam' and cursor+1<len(line):end=cursor+1;break
    first,last=line[index],line[end]
    date_rect=(first[0]-2,min(word[1] for word in line[index:end+1])-2,last[2]+2,max(word[3] for word in line[index:end+1])+2)
    break
 if not number_rect or not date_rect:return None
 def region(rect):
  x0,y0,x1,y1=rect
  x0=max(0,x0);y0=max(0,y0);x1=min(page_width,x1);y1=min(page_height,y1)
  return NumberRegion(x=x0/page_width*100,y=y0/page_height*100,width=(x1-x0)/page_width*100,height=(y1-y0)/page_height*100)
 try:
  result=NumberPlacement(number=region(number_rect),date=region(date_rect))
  return parse_number_placement(result.model_dump_json())
 except (ValueError,HTTPException):return None

def detect_outgoing_number_placement(source:Path)->NumberPlacement:
 with pymupdf.open(source) as doc:
  if not len(doc):raise HTTPException(422,'PDF không có trang để cấp số')
  page=doc[0];placement=_number_placement_from_words(page.get_text('words',sort=True),page.rect.width,page.rect.height)
  if placement:return placement
  pix=page.get_pixmap(matrix=pymupdf.Matrix(2,2),alpha=False)
 try:
  image=Image.open(io.BytesIO(pix.tobytes('png')))
  data=None
  for language in (os.getenv('OCR_LANGUAGES','vie+eng'),'eng'):
   try:data=pytesseract.image_to_data(image,lang=language,output_type=pytesseract.Output.DICT);break
   except pytesseract.TesseractError:continue
  if data:
   words=[]
   for index,text in enumerate(data['text']):
    if not text.strip():continue
    x,y,w,h=data['left'][index],data['top'][index],data['width'][index],data['height'][index]
    words.append((x,y,x+w,y+h,text,data['block_num'][index],data['line_num'][index],index))
   placement=_number_placement_from_words(words,image.width,image.height)
   if placement:return placement
 finally:image.close()
 raise HTTPException(422,'Không tự nhận diện được vị trí số và ngày trên trang đầu. Vui lòng chọn vùng thủ công')

def parse_number_placement(value:str)->NumberPlacement:
 try:
  placement=NumberPlacement.model_validate_json(value)
  for region in (placement.number,placement.date):
   if region.x+region.width>100 or region.y+region.height>100:raise ValueError('Vùng nằm ngoài trang')
  a,b=placement.number,placement.date
  if a.x<b.x+b.width and b.x<a.x+a.width and a.y<b.y+b.height and b.y<a.y+a.height:raise ValueError('Hai vùng bị chồng lên nhau')
  return placement
 except ValueError as exc:raise HTTPException(422,'Chọn hai vùng số và ngày hợp lệ, không chồng nhau') from exc

def render_outgoing_number(source:Path,number:int,issued:date,placement:NumberPlacement,preview=False)->bytes:
 with pymupdf.open(source) as doc:
  page=doc[0]
  regions=[]
  for region in (placement.number,placement.date):
   rect=pymupdf.Rect(region.x*page.rect.width/100,region.y*page.rect.height/100,(region.x+region.width)*page.rect.width/100,(region.y+region.height)*page.rect.height/100)
   regions.append(rect*page.derotation_matrix)
  # Remove only the selected fields, including pixels in scanned PDFs.
  for rect in regions:page.add_redact_annot(rect,fill=(1,1,1))
  page.apply_redactions(images=2,graphics=0,text=0)
  values=[str(number),f'ngày {issued.day:02d} tháng {issued.month:02d} năm {issued.year}']
  for index,(rect,value) in enumerate(zip(regions,values)):
   css='body {margin:0; font-family:serif; font-size:14pt; line-height:1; text-align:center;}' if index==0 else 'body {margin:0; font-family:serif; font-size:12pt;}'
   scale,_=page.insert_htmlbox(rect,html.escape(value),css=css,rotate=page.rotation,scale_low=12/14 if index==0 else 0.5)
   if scale<0:raise HTTPException(422,'Vùng đã chọn quá nhỏ, hãy kéo rộng hơn')
  if preview:return page.get_pixmap(matrix=pymupdf.Matrix(1.3,1.3)).tobytes('png')
  return doc.tobytes(garbage=4,deflate=True)

@app.post('/api/documents/{did}/outgoing-number-preview')
def outgoing_number_preview(did:int,placement:Optional[str]=Form(None),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or d.direction!='OUT':raise HTTPException(404,'Không tìm thấy văn bản')
 require_document_access(s,u,d);require_permission(s,u,'ASSIGN_NUMBER')
 if d.status!='READY_FOR_CLERK':raise HTTPException(409,'Văn bản chưa sẵn sàng cấp số')
 positions=parse_number_placement(placement) if placement else detect_outgoing_number_placement(Path(d.file_path));issued=d.issued_date or d.received_date or date.today()
 category='CÔNG VĂN ĐI' if d.doc_type in {'CÔNG VĂN','CÔNG VĂN NỘI BỘ'} else d.doc_type
 seq=s.query(Sequence).filter_by(year=issued.year,doc_type=category).first()
 number=d.number or (system_number_floor(s,issued.year,category)+1)
 return Response(render_outgoing_number(Path(d.file_path),number,issued,positions,preview=True),media_type='image/png')
def render_seal(source:Path,x_percent:int,y_percent:int,page_number:int,seal_type:str)->Path:
 doc=pymupdf.open(source);page=doc[max(0,min(len(doc)-1,page_number-1))];cx,cy=pdf_point(page,x_percent,y_percent);size=112;visual=pymupdf.Rect(cx-size/2,cy-size/2,cx+size/2,cy+size/2);outer=visual*page.derotation_matrix;inner=pymupdf.Rect(outer.x0+7,outer.y0+7,outer.x1-7,outer.y1-7);red=(.75,.04,.08);page.draw_oval(outer,color=red,width=2.5,overlay=True);page.draw_oval(inner,color=red,width=1.2,overlay=True);name='DAU TRUONG' if seal_type=='SCHOOL' else 'DAU VAN PHONG';page.insert_textbox(outer,f'\n{name}\nGIA LAP\nDH HUNG VUONG',fontname='helv',fontsize=7,color=red,align=1,rotate=page.rotation,overlay=True);target=STORE/f'{uuid.uuid4()}-sealed.pdf';doc.save(target,garbage=4,deflate=True);doc.close();return target
def append_opinion(source:Path,note:str,author:str,symbol:str,title:str)->Path:
 base=pymupdf.open(source);sheet=pymupdf.open();page=sheet.new_page(width=595,height=842);regular=Path('/System/Library/Fonts/Supplemental/Arial.ttf');bold=Path('/System/Library/Fonts/Supplemental/Arial Bold.ttf')
 if regular.exists():page.insert_font(fontname='arial',fontfile=str(regular))
 if bold.exists():page.insert_font(fontname='arial-bold',fontfile=str(bold))
 normal='arial' if regular.exists() else 'helv';strong='arial-bold' if bold.exists() else 'hebo';red=(.48,.04,.07);gray=(.25,.3,.36)
 page.insert_textbox(pymupdf.Rect(55,48,540,78),'TRƯỜNG ĐẠI HỌC HÙNG VƯƠNG TP. HỒ CHÍ MINH',fontname=strong,fontsize=10,color=gray,align=1)
 page.insert_textbox(pymupdf.Rect(55,92,540,135),'PHIẾU Ý KIẾN XỬ LÝ CÔNG VĂN ĐẾN',fontname=strong,fontsize=17,color=red,align=1)
 page.insert_textbox(pymupdf.Rect(65,145,530,190),f'Số/Ký hiệu: {symbol or "Chưa cấp số"}\nTrích yếu: {title}',fontname=normal,fontsize=9.5,color=gray,lineheight=1.25)
 box=pymupdf.Rect(65,220,530,610);page.draw_rect(box,color=(.58,.62,.68),width=1.2);page.insert_textbox(pymupdf.Rect(82,238,513,270),'Ý KIẾN CỦA CHÁNH VĂN PHÒNG',fontname=strong,fontsize=11,color=red,align=1)
 page.insert_htmlbox(pymupdf.Rect(88,282,507,505),'<div>'+html.escape(note).replace('\n','<br>')+'</div>',css='div {font-family: sans-serif; font-size: 12pt;}',scale_low=0)
 page.insert_textbox(pymupdf.Rect(300,510,500,535),f'TP. Hồ Chí Minh, {datetime.now().strftime("ngày %d/%m/%Y")}',fontname=normal,fontsize=9,color=gray,align=1)
 sign=pymupdf.Rect(330,540,485,600);page.draw_oval(sign,color=(.75,.04,.08),width=1.8);page.insert_textbox(sign,'\nĐÃ KÝ SỐ - GIẢ LẬP\n'+author,fontname=strong,fontsize=8,color=(.75,.04,.08),align=1)
 page.insert_textbox(pymupdf.Rect(65,650,530,720),'Phiếu được tạo và ký điện tử trên Hệ thống Quản lý văn bản. Nội dung được khóa sau khi xác nhận.',fontname=normal,fontsize=8.5,color=(.4,.44,.5),align=1)
 combined=pymupdf.open();combined.insert_pdf(sheet);combined.insert_pdf(base);target=STORE/f'{uuid.uuid4()}-office-opinion.pdf';combined.save(target,garbage=4,deflate=True);combined.close();base.close();sheet.close();return target
def ocr_pdf(path:Path)->str:
 doc=pymupdf.open(path);parts=[]
 for page in doc:
  native=page.get_text().strip()
  if len(native)>=80:parts.append(native);continue
  pix=page.get_pixmap(matrix=pymupdf.Matrix(2,2),alpha=False);img=Image.open(io.BytesIO(pix.tobytes('png')));parts.append(pytesseract.image_to_string(img,lang=os.getenv('OCR_LANGUAGES','vie+eng')))
 doc.close();return '\n'.join(parts)
@app.post('/api/documents/preview-upload')
async def preview_document_upload(file:UploadFile=File(...),u=Depends(current),s:Session=Depends(db)):
 if not any(has_permission(s,u,p) for p in ('CREATE','EDIT')):raise HTTPException(403,'Bạn không có quyền tạo hoặc sửa văn bản')
 return Response(await read_document_upload(file),media_type='application/pdf')

@app.post('/api/documents/analyze-title')
async def analyze_document_title(file:UploadFile=File(...),u=Depends(current),s:Session=Depends(db)):
 if not any(has_permission(s,u,p) for p in ('CREATE','EDIT')):raise HTTPException(403,'Bạn không có quyền tạo hoặc sửa văn bản')
 content=await read_document_upload(file,20*1024*1024)
 def analyze():
  parts=[]
  with pymupdf.open(stream=content,filetype='pdf') as doc:
   if doc.needs_pass:raise ValueError('PDF được bảo vệ bằng mật khẩu')
   for page in list(doc)[:3]:
    native=page.get_text().strip()
    if len(native)<80:
     pix=page.get_pixmap(matrix=pymupdf.Matrix(2,2),alpha=False)
     with Image.open(io.BytesIO(pix.tobytes('png'))) as img:
      native=pytesseract.image_to_string(img,lang=os.getenv('OCR_LANGUAGES','vie+eng'),timeout=30)
    parts.append(native)
  text='\n'.join(parts).strip()
  if not text:raise ValueError('Không đọc được nội dung PDF; vui lòng nhập trích yếu')
  title=extract_document_title(text)
  return {'title':title,'file_name':document_file_name(title)}
 try:return await run_in_threadpool(analyze)
 except ValueError as exc:raise HTTPException(422,str(exc))
 except Exception:raise HTTPException(503,'OCR/AI chưa sẵn sàng hoặc không đọc được PDF. Bạn có thể nhập trích yếu thủ công.')

def archive_file(d:Document,source:Optional[DocumentSource]=None):
 path=Path(source.file_path) if source else Path(d.file_path) if d.file_path else None
 if not path or not path.exists():return
 year=d.archive_year or (d.issued_date or d.received_date or d.created_at.date()).year;category=(d.doc_type if d.direction=='ARCHIVE' else ('CÔNG VĂN ĐẾN' if d.direction=='IN' else 'CÔNG VĂN ĐI') if d.doc_type=='CÔNG VĂN' else d.doc_type);folder=ARCHIVE_ROOT/str(year)/category
 folder.mkdir(parents=True,exist_ok=True);safe=''.join(c for c in (d.file_name or f'van-ban-{d.id}.pdf') if c not in '/\\:');target=folder/f'{d.id:06d}-{safe}'
 if path.resolve()!=target.resolve():shutil.copy2(path,target)
 if source:source.file_path=str(target);source.file_name=d.file_name or source.file_name
 else:d.file_path=str(target)
def incoming_source(s:Session,d:Document)->Optional[DocumentSource]:
 if d.direction!='IN' or not d.file_path:return None
 source=s.query(DocumentSource).filter_by(document_id=d.id).first()
 if source:
  if d.status in {'RECEIVED','NUMBERED'}:
   source.file_path=d.file_path;source.file_name=d.file_name or source.file_name
  return source
 source_path=Path(d.file_path)
 if not source_path.exists():return None
 opinion_statuses={'OFFICE_OPINION_COMPLETED','PENDING_BGH','BGH_APPROVED','BGH_REJECTED','FORWARDED','ARCHIVED'}
 if d.status in opinion_statuses:
  combined=pymupdf.open(source_path)
  if len(combined)<2:combined.close();return None
  original=pymupdf.open();original.insert_pdf(combined,from_page=1);target=STORE/f'{uuid.uuid4()}-incoming-original.pdf';original.save(target,garbage=4,deflate=True);original.close();combined.close();source_path=target
 source=DocumentSource(document_id=d.id,file_path=str(source_path),file_name=d.file_name or source_path.name);s.add(source);s.flush();return source
def migrate_opinion_sheets_to_note_only():
 """Regenerate active legacy opinion sheets that embedded workflow metadata."""
 with DB() as s:
  documents=s.query(Document).filter(Document.direction=='IN',Document.status.notin_({'ARCHIVED','FORWARDED','DELETED'})).all()
  for document in documents:
   reviews=s.query(IncomingReview).filter_by(document_id=document.id).order_by(IncomingReview.revision).all()
   if not reviews:continue
   source=incoming_source(s,document)
   source_path=resolve_source_file(source,document) if source else None
   if not source_path:continue
   for review in reviews:
    old_path=Path(review.pdf_path)
    if not old_path.exists():continue
    try:
     with pymupdf.open(old_path) as pdf:first_page=pdf[0].get_text() if len(pdf) else ''
    except Exception:continue
    if 'Vòng duyệt' not in first_page and 'Đơn vị nhận:' not in first_page:continue
    author=s.get(User,review.created_by)
    review.pdf_path=str(append_opinion(source_path,review.note,author.full_name if author else 'Chánh Văn phòng',document.symbol or '',document.title))
   document.file_path=reviews[-1].pdf_path
  s.commit()
migrate_opinion_sheets_to_note_only()
@app.post('/api/auth/login')
def login(x:Login,s:Session=Depends(db)):
 u=s.query(User).filter_by(username=x.username).first()
 if not u or not pwd.verify(x.password,u.password_hash):raise HTTPException(401,'Sai tài khoản hoặc mật khẩu')
 if not u.active:raise HTTPException(403,'Tài khoản đã bị khóa')
 return {'access_token':jwt.encode({'sub':str(u.id),'exp':datetime.utcnow()+timedelta(hours=12)},SECRET,algorithm='HS256'),'user':{'id':u.id,'full_name':u.full_name,'role':u.role,'department_id':u.department_id}}
@app.get('/api/dashboard')
def dashboard(u=Depends(current),s:Session=Depends(db)):
 require_office_module(u)
 a=[d for d in s.query(Document).filter(Document.status!='DELETED').all() if view(u,d,s)];return {'total':len(a),'incoming':sum(d.direction=='IN' for d in a),'outgoing':sum(d.direction in {'OUT','INTERNAL'} for d in a),'pending':sum(d.status.startswith('PENDING') for d in a),'archived':sum(d.status=='ARCHIVED' for d in a),'recent':[out(d) for d in sorted(a,key=lambda x:x.updated_at,reverse=True)[:6]]}

def meeting_report_data(start:date,end:date,u:User,s:Session):
 require_office_module(u)
 if end<start:raise HTTPException(422,'Ngày kết thúc phải từ ngày bắt đầu trở đi')
 if (end-start).days>3660:raise HTTPException(422,'Khoảng báo cáo tối đa 10 năm')
 documents=[]
 for document in s.query(Document).filter(Document.status!='DELETED').all():
  if not view(u,document,s):continue
  report_date=(date(document.archive_year,1,1) if document.direction=='ARCHIVE' and document.archive_year and not document.issued_date else (document.received_date if document.direction=='IN' else document.issued_date) or document.received_date or document.created_at.date())
  if start<=report_date<=end:documents.append((document,report_date))
 by_type={};by_status={}
 for document,_ in documents:
  by_type[document.doc_type]=by_type.get(document.doc_type,0)+1;by_status[document.status]=by_status.get(document.status,0)+1
 rows=[{'id':document.id,'date':report_date,'symbol':document.symbol,'title':document.title,'direction':document.direction,'doc_type':document.doc_type,'status':document.status,'issuing_agency':document.issuing_agency,'file_name':document.file_name,'uploaded_at':document.created_at} for document,report_date in sorted(documents,key=lambda item:(item[1],item[0].number or 0,item[0].id),reverse=True)]
 return {'start':start,'end':end,'total':len(documents),'incoming':sum(document.direction=='IN' for document,_ in documents),'outgoing':sum(document.direction in {'OUT','INTERNAL'} for document,_ in documents),'direct_archive':sum(document.direction=='ARCHIVE' for document,_ in documents),'by_type':[{'name':name,'count':count} for name,count in sorted(by_type.items(),key=lambda item:(-item[1],item[0]))],'by_status':[{'name':name,'count':count} for name,count in sorted(by_status.items(),key=lambda item:(-item[1],item[0]))],'rows':rows}

@app.get('/api/dashboard/meeting-report')
def meeting_report(start:date,end:date,u=Depends(current),s:Session=Depends(db)):
 return meeting_report_data(start,end,u,s)

MEETING_DIRECTIONS={'IN':'Văn bản đến','OUT':'Văn bản đi','INTERNAL':'Văn bản đi','ARCHIVE':'Tải trực tiếp'}
MEETING_STATUSES={'ARCHIVED':'Đã lưu trữ','RECEIVED':'Đã tiếp nhận','NUMBERED':'Đã cấp số','FORWARDED':'Đã chuyển đơn vị','DRAFT':'Bản nháp','RETURNED':'Đã trả lại','READY_FOR_CLERK':'Chờ Văn thư','OUT_NUMBERED':'Đã cấp số','SEALED':'Đã đóng dấu','EMAILED':'Đã phát hành','PENDING_ARCHIVE_REVIEW':'Chờ duyệt lưu trữ','ARCHIVE_RETURNED':'Đã trả lại'}
def _xlsx_column(index:int)->str:
 result=''
 while index:index,remainder=divmod(index-1,26);result=chr(65+remainder)+result
 return result
def _xlsx_cell(row:int,column:int,value,style:int=0,number=False)->str:
 ref=f'{_xlsx_column(column)}{row}'
 if number:return f'<c r="{ref}" s="{style}"><v>{int(value)}</v></c>'
 return f'<c r="{ref}" s="{style}" t="inlineStr"><is><t xml:space="preserve">{html.escape(str(value or ""))}</t></is></c>'
def render_meeting_xlsx(report:dict)->bytes:
 rows=[];merges=['A1:I1','A2:I2']
 rows.append('<row r="1" ht="28">'+_xlsx_cell(1,1,'BÁO CÁO GIAO BAN THÁNG',1)+'</row>')
 rows.append('<row r="2" ht="21">'+_xlsx_cell(2,1,f"Từ {report['start'].strftime('%d/%m/%Y')} đến {report['end'].strftime('%d/%m/%Y')}",2)+'</row>')
 metrics=[('Tổng văn bản',report['total']),('Văn bản đến',report['incoming']),('Văn bản đi',report['outgoing']),('Tải trực tiếp',report['direct_archive'])]
 label_cells=[];value_cells=[]
 for index,(label,value) in enumerate(metrics):
  column=index*2+1;merges.extend([f'{_xlsx_column(column)}4:{_xlsx_column(column+1)}4',f'{_xlsx_column(column)}5:{_xlsx_column(column+1)}5']);label_cells.append(_xlsx_cell(4,column,label,3));value_cells.append(_xlsx_cell(5,column,value,4,True))
 rows.extend(['<row r="4" ht="22">'+''.join(label_cells)+'</row>','<row r="5" ht="30">'+''.join(value_cells)+'</row>'])
 row_number=7;merges.append(f'A{row_number}:C{row_number}');rows.append(f'<row r="{row_number}" ht="22">'+_xlsx_cell(row_number,1,'CƠ CẤU THEO LOẠI',5)+'</row>');row_number+=1
 rows.append(f'<row r="{row_number}">'+_xlsx_cell(row_number,1,'Loại văn bản',5)+_xlsx_cell(row_number,2,'Số lượng',5)+'</row>');row_number+=1
 for item in report['by_type']:rows.append(f'<row r="{row_number}">'+_xlsx_cell(row_number,1,item['name'],6)+_xlsx_cell(row_number,2,item['count'],6,True)+'</row>');row_number+=1
 row_number+=1;header_row=row_number;headers=['Ngày','Số / ký hiệu','Trích yếu','Chiều văn bản','Loại văn bản','Trạng thái','Đơn vị','Tên file','Ngày tải lên'];rows.append(f'<row r="{row_number}" ht="28">'+''.join(_xlsx_cell(row_number,index,value,5) for index,value in enumerate(headers,1))+'</row>');row_number+=1
 for item in report['rows']:
  values=[item['date'].strftime('%d/%m/%Y'),item['symbol'] or 'Chưa cấp số',item['title'],MEETING_DIRECTIONS.get(item['direction'],item['direction']),item['doc_type'],MEETING_STATUSES.get(item['status'],item['status']),item['issuing_agency'],item['file_name'] or '',item['uploaded_at'].strftime('%d/%m/%Y %H:%M')]
  rows.append(f'<row r="{row_number}" ht="30">'+''.join(_xlsx_cell(row_number,index,value,6) for index,value in enumerate(values,1))+'</row>');row_number+=1
 worksheet='<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetViews><sheetView workbookViewId="0"><pane ySplit="'+str(header_row)+'" topLeftCell="A'+str(header_row+1)+'" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews><cols><col min="1" max="1" width="13" customWidth="1"/><col min="2" max="2" width="21" customWidth="1"/><col min="3" max="3" width="48" customWidth="1"/><col min="4" max="6" width="20" customWidth="1"/><col min="7" max="8" width="30" customWidth="1"/><col min="9" max="9" width="20" customWidth="1"/></cols><sheetData>'+''.join(rows)+'</sheetData><mergeCells count="'+str(len(merges))+'">'+''.join(f'<mergeCell ref="{ref}"/>' for ref in merges)+'</mergeCells><autoFilter ref="A'+str(header_row)+':I'+str(max(header_row,row_number-1))+'"/><pageMargins left="0.3" right="0.3" top="0.5" bottom="0.5" header="0.2" footer="0.2"/><pageSetup orientation="landscape" fitToWidth="1" fitToHeight="0"/></worksheet>'
 styles='<?xml version="1.0" encoding="UTF-8" standalone="yes"?><styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><fonts count="4"><font><sz val="11"/><name val="Arial"/></font><font><b/><color rgb="FFFFFFFF"/><sz val="18"/><name val="Arial"/></font><font><i/><color rgb="FF526071"/><sz val="11"/><name val="Arial"/></font><font><b/><color rgb="FFFFFFFF"/><sz val="11"/><name val="Arial"/></font></fonts><fills count="4"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor rgb="FF203571"/><bgColor indexed="64"/></patternFill></fill><fill><patternFill patternType="solid"><fgColor rgb="FFEAF0FA"/><bgColor indexed="64"/></patternFill></fill></fills><borders count="2"><border/><border><left style="thin"><color rgb="FFD9E0EA"/></left><right style="thin"><color rgb="FFD9E0EA"/></right><top style="thin"><color rgb="FFD9E0EA"/></top><bottom style="thin"><color rgb="FFD9E0EA"/></bottom></border></borders><cellXfs count="7"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/><xf numFmtId="0" fontId="1" fillId="2" borderId="0" applyAlignment="1"><alignment horizontal="center" vertical="center"/></xf><xf numFmtId="0" fontId="2" fillId="0" borderId="0" applyAlignment="1"><alignment horizontal="center"/></xf><xf numFmtId="0" fontId="0" fillId="3" borderId="1" applyAlignment="1"><alignment horizontal="center" vertical="center"/></xf><xf numFmtId="0" fontId="1" fillId="2" borderId="1" applyAlignment="1"><alignment horizontal="center" vertical="center"/></xf><xf numFmtId="0" fontId="3" fillId="2" borderId="1" applyAlignment="1"><alignment horizontal="center" vertical="center" wrapText="1"/></xf><xf numFmtId="0" fontId="0" fillId="0" borderId="1" applyAlignment="1"><alignment vertical="center" wrapText="1"/></xf></cellXfs></styleSheet>'
 stream=io.BytesIO()
 with zipfile.ZipFile(stream,'w',zipfile.ZIP_DEFLATED) as book:
  book.writestr('[Content_Types].xml','<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>')
  book.writestr('_rels/.rels','<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
  book.writestr('xl/workbook.xml','<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Báo cáo giao ban" sheetId="1" r:id="rId1"/></sheets></workbook>')
  book.writestr('xl/_rels/workbook.xml.rels','<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>')
  book.writestr('xl/styles.xml',styles);book.writestr('xl/worksheets/sheet1.xml',worksheet)
 return stream.getvalue()
REGISTER_SHEETS={'CÔNG VĂN ĐẾN':'CV Đến','CÔNG VĂN ĐI':'CV ĐI','QUYẾT ĐỊNH':'QĐ','THÔNG BÁO':'TB','KẾ HOẠCH':'KH','BÁO CÁO':'Báo cáo','NGHỊ QUYẾT HỘI ĐỒNG TRƯỜNG':'NQ-HĐT','THƯ MỜI':'Thư mời'}
REGISTER_TITLES={sheet:category for category,sheet in REGISTER_SHEETS.items()}
ROLE_NAMES={'ADMIN':'Quản trị hệ thống','BGH':'Ban Giám hiệu','OFFICE_HEAD':'Chánh Văn phòng','CLERK':'Văn thư','DEPARTMENT_HEAD':'Lãnh đạo đơn vị','DEPARTMENT':'Chuyên viên đơn vị'}
def document_register_category(document:Document)->str:
 if document.direction=='IN':return 'CÔNG VĂN ĐẾN'
 if document.doc_type in {'CÔNG VĂN','CÔNG VĂN ĐI','CÔNG VĂN NỘI BỘ'}:return 'CÔNG VĂN ĐI'
 return document.doc_type if document.doc_type in REGISTER_SHEETS else 'HỒ SƠ KHÁC'
def document_register_data(year:int,u:User,s:Session)->dict[str,list[list[str]]]:
 if year<1900 or year>date.today().year+1:raise HTTPException(422,'Năm sổ văn bản không hợp lệ')
 documents=[]
 for document in s.query(Document).filter(Document.status!='DELETED',Document.number.is_not(None)).all():
  document_year=document.archive_year or ((document.received_date if document.direction=='IN' else document.issued_date) or document.received_date or document.created_at.date()).year
  if document_year==year and view(u,document,s):documents.append(document)
 signatures={}
 for signature,user in s.query(DigitalSignature,User).join(User,DigitalSignature.user_id==User.id).order_by(DigitalSignature.id):signatures[signature.document_id]=(user.full_name,ROLE_NAMES.get(user.role,user.role))
 result={sheet:[] for sheet in REGISTER_TITLES}
 for document in sorted(documents,key=lambda item:(document_register_category(item),item.number or 0,item.id)):
  category=document_register_category(document);sheet=REGISTER_SHEETS.get(category)
  if not sheet:continue
  signed_by,position=signatures.get(document.id,('',''))
  issued=(document.issued_date or document.received_date or document.created_at.date()).strftime('%d/%m/%Y')
  title=(document.doc_type.title()+' - ' if document.doc_type not in {'CÔNG VĂN','CÔNG VĂN ĐẾN','CÔNG VĂN ĐI'} else '')+document.title
  if sheet=='CV Đến':
   review=latest_review(s,document.id);assignments=json.loads(review.assignments) if review else [];deadline=next((item.get('deadline') for item in assignments if item.get('deadline')),None)
   transfer=review.created_at.strftime('%d/%m/%Y') if review else '';units=', '.join(item.get('name','') for item in assignments)
   rows=result[sheet];rows.append([str(len(rows)+1),document.symbol or str(document.number),issued,title,signed_by,position,transfer,units,date.fromisoformat(deadline).strftime('%d/%m/%Y') if deadline else '',MEETING_STATUSES.get(document.status,document.status),document.summary or ''])
  else:
   rows=result[sheet];rows.append([str(len(rows)+1),document.symbol or str(document.number),issued,title,signed_by,position,document.summary or MEETING_STATUSES.get(document.status,document.status)])
 return result
def render_document_register_xlsx(year:int,data:dict[str,list[list[str]]])->bytes:
 if not REGISTER_TEMPLATE.exists():raise HTTPException(500,'Chưa cấu hình file mẫu sổ đăng ký văn bản')
 main_ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';office_rel='http://schemas.openxmlformats.org/officeDocument/2006/relationships';package_rel='http://schemas.openxmlformats.org/package/2006/relationships'
 q=lambda tag:f'{{{main_ns}}}{tag}'
 with zipfile.ZipFile(REGISTER_TEMPLATE,'r') as source:
  workbook=ET.fromstring(source.read('xl/workbook.xml'));relations=ET.fromstring(source.read('xl/_rels/workbook.xml.rels'));targets={node.attrib['Id']:node.attrib['Target'] for node in relations}
  sheet_paths={node.attrib['name']:'xl/'+targets[node.attrib[f'{{{office_rel}}}id']].lstrip('/') for node in workbook.find(q('sheets'))}
  replacements={}
  for sheet_name,rows in data.items():
   if sheet_name not in REGISTER_TITLES or sheet_name=="Mẫu":continue
   path=sheet_paths.get(sheet_name)
   if not path:continue
   worksheet=ET.fromstring(source.read(path));sheet_data=worksheet.find(q('sheetData'));max_columns=11 if sheet_name=='CV Đến' else 7
   row_nodes={int(node.attrib['r']):node for node in sheet_data.findall(q('row'))}
   for row_number,node in row_nodes.items():
    if row_number<8:continue
    for cell in node.findall(q('c')):
     match=re.match(r'([A-Z]+)',cell.attrib.get('r',''))
     if not match:continue
     column=sum((ord(char)-64)*26**index for index,char in enumerate(reversed(match.group(1))))
     if column<=max_columns:
      for child in list(cell):cell.remove(child)
      cell.attrib.pop('t',None)
   for offset,values in enumerate(rows,8):
    node=row_nodes.get(offset)
    if node is None:node=ET.SubElement(sheet_data,q('row'),{'r':str(offset)});row_nodes[offset]=node
    cells={cell.attrib.get('r'):cell for cell in node.findall(q('c'))}
    for column,value in enumerate(values,1):
     ref=f'{_xlsx_column(column)}{offset}';cell=cells.get(ref)
     if cell is None:cell=ET.SubElement(node,q('c'),{'r':ref,'s':'38'})
     for child in list(cell):cell.remove(child)
     cell.set('t','inlineStr');inline=ET.SubElement(cell,q('is'));text=ET.SubElement(inline,q('t'));text.text=str(value or '')
   title_cell=next((cell for cell in row_nodes.get(5,[]) if cell.attrib.get('r')=='A5'),None)
   if title_cell is not None:
    for child in list(title_cell):title_cell.remove(child)
    title_cell.set('t','inlineStr');inline=ET.SubElement(title_cell,q('is'));text=ET.SubElement(inline,q('t'));text.text=f'ĐĂNG KÝ SỐ VĂN BẢN: {REGISTER_TITLES[sheet_name]} - NĂM {year}'
   replacements[path]=ET.tostring(worksheet,encoding='utf-8',xml_declaration=True)
  stream=io.BytesIO()
  with zipfile.ZipFile(stream,'w',zipfile.ZIP_DEFLATED) as target:
   for item in source.infolist():target.writestr(item,replacements.get(item.filename,source.read(item.filename)))
 return stream.getvalue()
@app.get('/api/register/export')
def export_document_register(year:int,u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'ARCHIVE','Chỉ Văn thư hoặc Admin được xuất sổ quản lý số hiệu')
 content=render_document_register_xlsx(year,document_register_data(year,u,s))
 return Response(content,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="so-dang-ky-so-van-ban-{year}.xlsx"'})
def render_meeting_pdf(report:dict)->bytes:
 pdf=pymupdf.open();detail=report['rows'];chunks=[detail[:7]];remaining=detail[7:];chunks.extend(remaining[index:index+14] for index in range(0,len(remaining),14))
 type_summary=' · '.join(f"{html.escape(item['name'])}: <b>{item['count']}</b>" for item in report['by_type']) or 'Không có văn bản'
 for page_index,items in enumerate(chunks,1):
  page=pdf.new_page(width=842,height=595);rows=''.join('<tr><td>'+item['date'].strftime('%d/%m/%Y')+'</td><td>'+html.escape(item['symbol'] or 'Chưa cấp số')+'</td><td>'+html.escape(item['title'])+'</td><td>'+html.escape(MEETING_DIRECTIONS.get(item['direction'],item['direction']))+'</td><td>'+html.escape(item['doc_type'])+'</td><td>'+html.escape(MEETING_STATUSES.get(item['status'],item['status']))+'</td></tr>' for item in items)
  heading=('<div class="school">TRƯỜNG ĐẠI HỌC HÙNG VƯƠNG TP. HỒ CHÍ MINH</div><h1>BÁO CÁO GIAO BAN THÁNG</h1><div class="period">Từ '+report['start'].strftime('%d/%m/%Y')+' đến '+report['end'].strftime('%d/%m/%Y')+'</div><table class="metrics"><tr><td>TỔNG VĂN BẢN<b>'+str(report['total'])+'</b></td><td>VĂN BẢN ĐẾN<b>'+str(report['incoming'])+'</b></td><td>VĂN BẢN ĐI<b>'+str(report['outgoing'])+'</b></td><td>TẢI TRỰC TIẾP<b>'+str(report['direct_archive'])+'</b></td></tr></table><div class="types"><b>Cơ cấu theo loại:</b> '+type_summary+'</div>') if page_index==1 else '<h2>BÁO CÁO GIAO BAN THÁNG — CHI TIẾT (tiếp)</h2>'
  columns='<colgroup><col width="9%"><col width="16%"><col width="35%"><col width="12%"><col width="14%"><col width="14%"></colgroup>'
  content=heading+'<table class="details">'+columns+'<thead><tr><th>Ngày</th><th>Số / ký hiệu</th><th>Trích yếu</th><th>Chiều</th><th>Loại</th><th>Trạng thái</th></tr></thead><tbody>'+rows+'</tbody></table><div class="footer">Trang '+str(page_index)+'/'+str(len(chunks))+' · Xuất từ Hệ thống Quản lý văn bản DHV</div>'
  css='body{font-family:sans-serif;color:#26364a;font-size:8pt} .school{text-align:center;color:#203571;font-weight:bold;font-size:9pt} h1{text-align:center;color:#791019;font-size:20pt;margin:8px 0 2px} h2{text-align:center;color:#203571;font-size:13pt}.period{text-align:center;color:#64748b;margin-bottom:12px}.metrics,.details{border-collapse:collapse;width:100%}.details{table-layout:fixed}.metrics td{background:#203571;color:white;text-align:center;padding:8px;border-right:2px solid white}.metrics b{display:block;font-size:18pt;margin-top:3px}.types{background:#edf2fb;padding:8px;margin:8px 0}.details th{background:#791019;color:white;padding:6px;border:1px solid #b9c1cc}.details td{padding:5px;border:1px solid #d9dfe7;vertical-align:top;overflow-wrap:break-word}.footer{text-align:right;color:#7b8490;margin-top:8px;font-size:7pt}'
  spare,_=page.insert_htmlbox(pymupdf.Rect(28,22,814,574),content,css=css,scale_low=.45)
  if spare<0:raise HTTPException(500,'Không thể dàn trang báo cáo PDF')
 data=pdf.tobytes(garbage=4,deflate=True);pdf.close();return data

@app.get('/api/dashboard/meeting-report/export/{file_type}')
def export_meeting_report(file_type:Literal['xlsx','pdf'],start:date,end:date,u=Depends(current),s:Session=Depends(db)):
 report=meeting_report_data(start,end,u,s);name=f'bao-cao-giao-ban-{start.isoformat()}-{end.isoformat()}.{file_type}'
 if file_type=='xlsx':content=render_meeting_xlsx(report);media='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
 else:content=render_meeting_pdf(report);media='application/pdf'
 return Response(content,media_type=media,headers={'Content-Disposition':f'attachment; filename="{name}"'})
@app.get('/api/documents')
def docs(direction:Optional[str]=None,q:Optional[str]=None,status:Optional[str]=None,archive_workspace:bool=False,u=Depends(current),s:Session=Depends(db)):
 if archive_workspace or status=='ARCHIVED' or direction=='ARCHIVE':require_office_module(u)
 x=s.query(Document).filter(Document.status!='DELETED')
 if direction=='OUT':x=x.filter(Document.direction.in_(['OUT','INTERNAL']))
 elif direction:x=x.filter_by(direction=direction)
 if archive_workspace:x=x.filter(Document.status.in_({'ARCHIVED','PENDING_ARCHIVE_REVIEW','ARCHIVE_RETURNED'}))
 elif status:x=x.filter_by(status=status)
 if q:x=x.filter(Document.search_text.contains(normalize_search(q)))
 return [out(d) for d in x.order_by(Document.updated_at.desc()) if view(u,d,s)]

@app.post('/api/archive/uploads')
async def upload_archive_document(doc_type:str=Form('HỒ SƠ KHÁC'),folder:Optional[str]=Form(None),archive_year:Optional[int]=Form(None),title:str=Form(''),issued_date:Optional[str]=Form(None),file:UploadFile=File(...),u=Depends(current),s:Session=Depends(db)):
 require_office_module(u)
 require_permission(s,u,'CREATE','Bạn không có quyền gửi hồ sơ lưu trữ')
 archive_folder=(folder or doc_type).strip()
 if archive_folder=='CÔNG VĂN':archive_folder='CÔNG VĂN ĐI'
 if archive_folder not in ARCHIVE_FOLDERS:raise HTTPException(422,'Thư mục lưu trữ không hợp lệ')
 chosen_year=archive_year or (date.fromisoformat(issued_date).year if issued_date else date.today().year)
 if chosen_year<1900 or chosen_year>date.today().year+1:raise HTTPException(422,'Năm lưu trữ không hợp lệ')
 content=await read_document_upload(file)
 path=STORE/f'{uuid.uuid4()}-archive-upload.pdf';path.write_bytes(content)
 dep=s.get(Department,u.department_id) if u.department_id else None
 direct=has_permission(s,u,'ARCHIVE');document=Document(direction='ARCHIVE',doc_type=archive_folder,title=title.strip() or Path(file.filename).stem,summary='Tải trực tiếp lên kho lưu trữ',issuing_agency=dep.name if dep else u.full_name,issued_date=date.fromisoformat(issued_date) if issued_date else None,received_date=date.today(),archive_year=chosen_year,department_id=u.department_id,owner_id=u.id,status='ARCHIVED' if direct else 'PENDING_ARCHIVE_REVIEW',file_name=Path(file.filename).with_suffix('.pdf').name,file_path=str(path))
 s.add(document);s.flush()
 if direct:
  document.number=allocate_archive_number(s,chosen_year,archive_folder);document.symbol=number_symbol(archive_folder,document.number,date(chosen_year,1,1));document.ocr_text=ocr_pdf(path);archive_file(document);action='ARCHIVE_DIRECT_UPLOAD'
 else:action='ARCHIVE_UPLOAD'
 refresh_search(document);log(s,document,u,action,json.dumps({'file_name':document.file_name,'folder':archive_folder,'archive_year':chosen_year,'uploaded_at':document.created_at.isoformat()},ensure_ascii=False));s.commit();return out(document)

@app.post('/api/archive/uploads/{did}/resubmit')
async def resubmit_archive_document(did:int,title:str=Form(''),issued_date:Optional[str]=Form(None),file:UploadFile=File(...),u=Depends(current),s:Session=Depends(db)):
 require_office_module(u)
 document=s.get(Document,did)
 if not document or document.direction!='ARCHIVE' or document.owner_id!=u.id:raise HTTPException(404,'Không tìm thấy hồ sơ tải lên')
 if document.status!='ARCHIVE_RETURNED':raise HTTPException(409,'Hồ sơ không ở trạng thái cần nộp lại')
 content=await read_document_upload(file)
 path=STORE/f'{uuid.uuid4()}-archive-resubmit.pdf';path.write_bytes(content)
 old=Path(document.file_path) if document.file_path else None
 if old and old.is_file() and old.parent.resolve()==STORE.resolve():old.unlink()
 document.file_path=str(path);document.file_name=Path(file.filename).with_suffix('.pdf').name;document.title=title.strip() or Path(file.filename).stem;document.issued_date=date.fromisoformat(issued_date) if issued_date else document.issued_date;document.revision+=1;document.status='PENDING_ARCHIVE_REVIEW';refresh_search(document);log(s,document,u,'ARCHIVE_RESUBMIT',f'Phiên bản {document.revision}: {document.file_name}');s.commit();return out(document)

@app.post('/api/archive/uploads/{did}/review')
def review_archive_document(did:int,x:ArchiveReviewDecision,u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'ARCHIVE','Bạn không có quyền duyệt hồ sơ lưu trữ')
 document=s.get(Document,did)
 if not document or document.direction!='ARCHIVE':raise HTTPException(404,'Không tìm thấy hồ sơ tải lên')
 if document.status!='PENDING_ARCHIVE_REVIEW':raise HTTPException(409,'Hồ sơ đã được xử lý')
 if not x.note.strip():raise HTTPException(422,'Ý kiến duyệt hoặc lý do trả lại là bắt buộc')
 if x.decision=='RETURN':
  document.status='ARCHIVE_RETURNED';log(s,document,u,'ARCHIVE_RETURN',json.dumps({'revision':document.revision,'reason':x.note.strip()},ensure_ascii=False));s.commit();return out(document)
 issued=document.issued_date or document.received_date or date.today();year=document.archive_year or issued.year;document.number=allocate_archive_number(s,year,document.doc_type);document.archive_year=year;document.issued_date=issued;document.symbol=number_symbol(document.doc_type,document.number,date(year,1,1));document.file_name=document_file_name(document.title,document.symbol)
 document.ocr_text=ocr_pdf(Path(document.file_path));archive_file(document);document.status='ARCHIVED';refresh_search(document);log(s,document,u,'ARCHIVE_APPROVE',json.dumps({'revision':document.revision,'note':x.note.strip(),'symbol':document.symbol},ensure_ascii=False));s.commit();return out(document)
@app.post('/api/documents')
async def create(direction:str=Form(...),doc_type:str=Form('CÔNG VĂN'),title:str=Form(...),summary:str=Form(''),issuing_agency:str=Form(''),issued_date:Optional[str]=Form(None),department_id:Optional[int]=Form(None),priority:str=Form('Bình thường'),file:Optional[UploadFile]=File(None),u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'CREATE','Bạn không có quyền tạo văn bản')
 if direction not in {'IN','OUT','INTERNAL'}:raise HTTPException(400,'Luồng văn bản không hợp lệ')
 if u.role in {'DEPARTMENT','DEPARTMENT_HEAD'} and direction!='OUT':raise HTTPException(403,'Đơn vị chỉ được tạo công văn đi / nội bộ trong tab Công văn đi')
 if direction=='IN':doc_type='CÔNG VĂN'
 if direction=='OUT' and doc_type not in OUTGOING_TYPES:raise HTTPException(400,'Loại công văn không hợp lệ')
 if direction=='OUT' and doc_type in {*DOCUMENT_TYPE_CODES,'CÔNG VĂN NỘI BỘ'} and issuing_agency.strip() not in INTERNAL_UNITS:raise HTTPException(400,'Chọn đơn vị tiếp nhận trong trường')
 if direction=='INTERNAL' and doc_type not in DOCUMENT_TYPE_CODES:raise HTTPException(400,'Loại văn bản nội bộ không hợp lệ')
 if direction in {'IN','OUT','INTERNAL'}:
  if direction!='OUT' and not issued_date:raise HTTPException(422,'Ngày ban hành là bắt buộc')
  if not issuing_agency.strip():raise HTTPException(422,'Đơn vị phát hành là bắt buộc' if direction=='INTERNAL' else 'Đơn vị ban hành là bắt buộc' if direction=='IN' else 'Đơn vị tiếp nhận là bắt buộc')
  if direction=='INTERNAL' and issuing_agency.strip() not in INTERNAL_UNITS:raise HTTPException(400,'Đơn vị phát hành nội bộ không hợp lệ')
  if not file or not file.filename:raise HTTPException(422,'Bản scan / văn bản là bắt buộc')
  office=s.query(Department).filter_by(code='VP').first()
  if not office:raise HTTPException(500,'Chưa cấu hình đơn vị Văn phòng trường')
  department_id=u.department_id if direction=='OUT' else office.id
 if direction=='OUT':
  if u.department_id is None:raise HTTPException(422,'Tài khoản cần được gán đơn vị soạn thảo')
 d=Document(direction=direction,doc_type=doc_type,title=title,summary=summary,issuing_agency=issuing_agency,issued_date=date.fromisoformat(issued_date) if issued_date else None,received_date=date.today() if direction=='IN' else None,department_id=department_id,priority=priority,owner_id=u.id,status='RECEIVED' if direction=='IN' else 'DRAFT')
 if file and file.filename:
  p=STORE/f'{uuid.uuid4()}.pdf';p.write_bytes(await read_document_upload(file));d.file_name=document_file_name(title);d.file_path=str(p)
  try:
   if p.suffix.lower()=='.pdf':d.ocr_text='\n'.join(x.get_text() for x in pymupdf.open(p))
  except Exception:d.ocr_text=''
 s.add(d);s.flush();refresh_search(d);log(s,d,u,'CREATE');s.commit();return out(d)
@app.post('/api/documents/{did}/outgoing-submit')
def outgoing_submit(did:int,bgh_action:str=Form(...),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or d.direction!='OUT':raise HTTPException(404,'Không tìm thấy công văn đi')
 require_permission(s,u,'EDIT','Bạn không có quyền trình công văn đi')
 if not view(u,d,s):raise HTTPException(404,'Không tìm thấy công văn đi')
 if not d.file_path or not d.issuing_agency.strip():raise HTTPException(422,'Văn bản chưa đủ thông tin bắt buộc')
 if bgh_action=='APPROVE':
  if d.status not in {'UNIT_SIGNED','OFFICE_SIGNED'}:raise HTTPException(409,'Lãnh đạo đơn vị phải ký trước khi gửi BGH duyệt')
  d.status='PENDING_OUT_BGH_APPROVAL';action='OUTGOING_SUBMIT_BGH_APPROVAL'
 elif bgh_action=='SIGN':
  if d.status not in {'DRAFT','RETURNED'}:raise HTTPException(409,'Chỉ bản chưa ký mới được gửi BGH ký số')
  d.status='PENDING_OUT_BGH_SIGN';action='OUTGOING_SUBMIT_BGH_SIGN'
 else:raise HTTPException(422,'Hình thức trình BGH không hợp lệ')
 log(s,d,u,action,f'Phiên bản {d.revision}');s.commit();return out(d)

@app.post('/api/documents/{did}/outgoing-clerk-approve')
def outgoing_clerk_approve(did:int,u=Depends(current),s:Session=Depends(db)):
 raise HTTPException(409,'Luồng đã thay đổi: văn bản được gửi thẳng BGH, sau khi duyệt sẽ chuyển về Văn thư')

@app.post('/api/documents/{did}/outgoing-bgh-approve')
def outgoing_bgh_approve(did:int,note:str=Form(...),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or d.direction!='OUT':raise HTTPException(404,'Không tìm thấy công văn đi')
 require_document_access(s,u,d)
 if u.role not in {'ADMIN','BGH'}:raise HTTPException(403,'Chỉ BGH được duyệt văn bản')
 if d.status!='PENDING_OUT_BGH_APPROVAL':raise HTTPException(409,'Văn bản không ở bước chờ BGH duyệt')
 if not d.signed_personal:raise HTTPException(409,'Văn bản chưa có chữ ký số của lãnh đạo đơn vị')
 if not note.strip():raise HTTPException(422,'Ý kiến duyệt là bắt buộc')
 d.received_date=date.today();d.status='READY_FOR_CLERK';log(s,d,u,'OUTGOING_BGH_APPROVE',json.dumps({'decision':'APPROVE','revision':d.revision,'note':note.strip()},ensure_ascii=False));s.commit();return out(d)
@app.post('/api/documents/{did}/outgoing-return')
def outgoing_return(did:int,note:str=Form(...),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or d.direction!='OUT':raise HTTPException(404,'Không tìm thấy công văn đi')
 require_document_access(s,u,d)
 if d.status in {'PENDING_OUT_BGH_APPROVAL','PENDING_OUT_BGH_SIGN'}:
  if u.role not in {'ADMIN','BGH'}:raise HTTPException(403,'Chỉ BGH được trả lại văn bản đã trình')
  actor='BGH'
 else:raise HTTPException(409,'Văn bản không ở bước có thể trả lại')
 if not note.strip():raise HTTPException(422,'Lý do trả lại là bắt buộc')
 source=s.query(DocumentSource).filter_by(document_id=d.id).first()
 if source:d.file_path=source.file_path
 old_revision=d.revision;d.revision+=1;d.signed_personal=False;d.signature_hash=None;d.status='RETURNED'
 log(s,d,u,'OUTGOING_RETURN',json.dumps({'actor':actor,'from_revision':old_revision,'to_revision':d.revision,'reason':note.strip()},ensure_ascii=False));s.commit();return out(d)
@app.post('/api/documents/{did}/outgoing-number')
def outgoing_number(did:int,placement:Optional[str]=Form(None),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or d.direction!='OUT':raise HTTPException(404,'Không tìm thấy công văn đi')
 require_document_access(s,u,d)
 require_permission(s,u,'ASSIGN_NUMBER','Bạn không có quyền cấp số')
 if d.status!='READY_FOR_CLERK' or not d.signed_personal:raise HTTPException(409,'Văn bản phải hoàn tất bước ký và duyệt trước khi cấp số')
 positions=parse_number_placement(placement) if placement else detect_outgoing_number_placement(Path(d.file_path))
 # Claim the transition before reading the sequence; SQLite serializes concurrent writers.
 claimed=s.execute(update(Document).where(Document.id==did,Document.status=='READY_FOR_CLERK').values(status='OUT_NUMBERING')).rowcount
 if not claimed:s.rollback();raise HTTPException(409,'Văn bản đang được cấp số hoặc đã có số')
 issued=d.issued_date or d.received_date or date.today()
 year=issued.year;sequence_type='CÔNG VĂN ĐI' if d.doc_type in {'CÔNG VĂN','CÔNG VĂN NỘI BỘ'} else d.doc_type
 if not d.number:d.number=allocate_archive_number(s,year,sequence_type)
 d.archive_year=year
 d.issued_date=issued;d.symbol=number_symbol(d.doc_type if d.doc_type!='CÔNG VĂN NỘI BỘ' else 'CÔNG VĂN',d.number,issued,'OUT');d.file_name=document_file_name(d.title,d.symbol);d.status='OUT_NUMBERED'
 try:
  content=render_outgoing_number(Path(d.file_path),d.number,issued,positions)
  target=STORE/f'{uuid.uuid4()}-numbered.pdf';target.write_bytes(content);d.file_path=str(target)
  refresh_search(d);log(s,d,u,'OUTGOING_NUMBER',json.dumps({'symbol':d.symbol,'placement':positions.model_dump(),'placement_mode':'MANUAL' if placement else 'AUTO'},ensure_ascii=False));s.commit()
 except Exception:
  s.rollback();raise
 return out(d)
@app.post('/api/documents/{did}/outgoing-email')
def outgoing_email(did:int,emails:str=Form(...),subject:str=Form(...),content_html:str=Form(...),signature_html:str=Form(...),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or d.direction!='OUT':raise HTTPException(404,'Không tìm thấy công văn đi')
 require_document_access(s,u,d);require_permission(s,u,'EMAIL','Bạn không có quyền phát hành văn bản')
 if d.status!='SEALED':raise HTTPException(409,'Văn bản phải được đóng dấu trước khi gửi mail')
 recipients=[x.strip() for x in re.split(r'[,;\n]',emails) if x.strip()]
 if not recipients:raise HTTPException(422,'Phải nhập ít nhất một email nhận')
 return gmail_draft(d,recipients,subject,content_html,signature_html)
@app.post('/api/documents/{did}/internal-submit')
def internal_submit(did:int,u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or d.direction!='INTERNAL':raise HTTPException(404,'Không tìm thấy văn bản nội bộ')
 require_document_access(s,u,d);require_permission(s,u,'EDIT','Chỉ Văn thư được trình văn bản nội bộ')
 if d.status!='INTERNAL_NUMBERED':raise HTTPException(409,'Văn bản nội bộ phải được cấp số trước khi trình duyệt')
 d.status='PENDING_INTERNAL_APPROVAL';log(s,d,u,'INTERNAL_SUBMIT_OFFICE');s.commit();return out(d)
@app.post('/api/documents/{did}/internal-approve')
def internal_approve(did:int,note:str=Form(''),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or d.direction!='INTERNAL':raise HTTPException(404,'Không tìm thấy văn bản nội bộ')
 require_document_access(s,u,d);require_permission(s,u,'INTERNAL_APPROVE','Chỉ Chánh Văn phòng được phê duyệt văn bản nội bộ')
 if d.status!='PENDING_INTERNAL_APPROVAL':raise HTTPException(409,'Văn bản nội bộ chưa được trình Chánh Văn phòng')
 d.status='INTERNAL_APPROVED';log(s,d,u,'INTERNAL_APPROVED',note.strip());s.commit();return out(d)
@app.post('/api/documents/{did}/internal-email')
def internal_email(did:int,emails:str=Form(...),subject:str=Form(...),content_html:str=Form(...),signature_html:str=Form(...),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or d.direction!='INTERNAL':raise HTTPException(404,'Không tìm thấy công văn nội bộ')
 require_document_access(s,u,d);require_permission(s,u,'EMAIL','Bạn không có quyền phát hành văn bản')
 if d.status!='INTERNAL_SEALED' or not d.sealed:raise HTTPException(409,'Văn bản nội bộ phải được đóng dấu trước khi phát hành')
 recipients=[x.strip() for x in re.split(r'[,;\n]',emails) if x.strip()]
 if not recipients:raise HTTPException(422,'Phải nhập ít nhất một email nhận')
 return gmail_draft(d,recipients,subject,content_html,signature_html)
@app.post('/api/documents/{did}/incoming-number')
def incoming_number(did:int,year:int=Form(default=date.today().year),number:Optional[int]=Form(None),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or d.direction!='IN':raise HTTPException(404,'Không tìm thấy công văn đến')
 require_document_access(s,u,d)
 require_permission(s,u,'ASSIGN_NUMBER','Bạn không có quyền vào số')
 if d.status!='RECEIVED':raise HTTPException(409,'Chỉ công văn vừa tiếp nhận mới được vào số')
 lock_review(s)
 floor=system_number_floor(s,year,'CÔNG VĂN ĐẾN')
 if number is not None and number!=floor+1:raise HTTPException(409,f'Số tiếp theo của hệ thống là {floor+1}')
 chosen=allocate_archive_number(s,year,'CÔNG VĂN ĐẾN')
 s.add(IncomingNumber(document_id=d.id,year=year,doc_type=d.doc_type,number=chosen));d.number=chosen;d.archive_year=year;d.symbol=number_symbol(d.doc_type,chosen,date(year,1,1),'IN');d.file_name=document_file_name(d.title,d.symbol) if d.file_name else None;d.status='NUMBERED';refresh_search(d);log(s,d,u,'INCOMING_NUMBER',d.symbol);s.commit();return out(d)
@app.post('/api/documents/{did}/edit-received')
async def edit_received(did:int,title:str=Form(...),doc_type:str=Form('CÔNG VĂN'),summary:str=Form(''),issuing_agency:str=Form(''),issued_date:Optional[str]=Form(None),file_name:Optional[str]=Form(None),file:Optional[UploadFile]=File(None),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or d.direction!='IN':raise HTTPException(404,'Không tìm thấy công văn đến')
 require_document_access(s,u,d)
 if d.status not in {'RECEIVED','NUMBERED'}:raise HTTPException(409,'Chỉ được chỉnh sửa công văn đến khi vừa tiếp nhận hoặc đã cấp số, trước khi có ý kiến Văn phòng')
 require_permission(s,u,'EDIT','Chỉ Văn thư được chỉnh sửa công văn đến')
 doc_type='CÔNG VĂN'
 if not issued_date:raise HTTPException(422,'Ngày ban hành là bắt buộc')
 if not issuing_agency.strip():raise HTTPException(422,'Đơn vị ban hành là bắt buộc')
 if not d.file_path and (not file or not file.filename):raise HTTPException(422,'Bản scan / văn bản là bắt buộc')
 d.title=title.strip();d.doc_type=doc_type;d.summary=summary;d.issuing_agency=issuing_agency;d.issued_date=date.fromisoformat(issued_date) if issued_date else None
 d.file_name=document_file_name(d.title,d.symbol) if d.file_name else None
 if file and file.filename:
  content=await read_document_upload(file);p=STORE/f'{uuid.uuid4()}.pdf';p.write_bytes(content);d.file_name=document_file_name(d.title,d.symbol);d.file_path=str(p)
  d.ocr_text=''
  source=s.query(DocumentSource).filter_by(document_id=d.id).first()
  if source:source.file_path=d.file_path;source.file_name=d.file_name
 refresh_search(d);log(s,d,u,'EDIT_RECEIVED',f'Tên file: {d.file_name or "không có"}');s.commit();return out(d)
def reservation_department(s,u,department_id=None):
 require_permission(s,u,'RESERVE_NUMBER','Bạn không có quyền xin số trước')
 selected=department_id if u.role in {'ADMIN','CLERK','OFFICE_HEAD'} else u.department_id
 if selected is None:raise HTTPException(422,'Vui lòng chọn đơn vị yêu cầu')
 department=s.get(Department,selected)
 if not department:raise HTTPException(422,'Tài khoản chưa được gán đơn vị hợp lệ')
 return department

@app.post('/api/reserve-number')
def reserve_number(direction:str=Form(...),doc_type:str=Form('CÔNG VĂN'),title:str=Form(...),issuing_agency:str=Form(''),reason:str=Form(...),issued_date:Optional[str]=Form(None),department_id:Optional[int]=Form(None),request_id:Optional[int]=Form(None),u=Depends(current),s:Session=Depends(db)):
 try:issued_date=date.fromisoformat(issued_date.strip()) if issued_date and issued_date.strip() else None
 except ValueError:raise HTTPException(422,'Ngày phát hành không hợp lệ')
 department=reservation_department(s,u,department_id)
 issuing_agency=department.name
 if direction not in {'OUT','INTERNAL'}:raise HTTPException(400,'Luồng xin số không hợp lệ')
 if direction=='OUT' and doc_type not in OUTGOING_TYPES:raise HTTPException(400,'Loại văn bản không hợp lệ')
 if direction=='INTERNAL' and doc_type not in DOCUMENT_TYPE_CODES:raise HTTPException(400,'Loại văn bản nội bộ không hợp lệ')
 if not title.strip() or not reason.strip():raise HTTPException(422,'Trích yếu dự kiến và lý do xin số là bắt buộc')
 if request_id:
  d=s.get(Document,request_id)
  if not d or d.owner_id!=u.id:raise HTTPException(404,'Không tìm thấy yêu cầu xin số')
  if d.status!='RESERVATION_REJECTED':raise HTTPException(409,'Chỉ được chỉnh yêu cầu đã bị từ chối')
  d.revision+=1
 else:
  d=Document(owner_id=u.id,reserved=True,archive_year=date.today().year)
  s.add(d)
 d.direction=direction;d.doc_type=doc_type;d.title=title.strip();d.summary=f'Lý do xin số trước: {reason.strip()}';d.issuing_agency=department.name;d.department_id=department.id;d.issued_date=issued_date;d.status='RESERVATION_PENDING'
 s.flush();refresh_search(d);log(s,d,u,'RESERVATION_RESUBMIT' if request_id else 'RESERVATION_SUBMIT',reason.strip());s.commit();return out(d)

class ReservationDecision(BaseModel):
 decision:Literal['APPROVE','REJECT','DELETE']
 revision:int
 note:str=''

@app.post('/api/documents/{did}/reservation-decision')
def reservation_decision(did:int,x:ReservationDecision,u=Depends(current),s:Session=Depends(db)):
 try:
  lock_review(s)
  d=s.get(Document,did)
  if not d:raise HTTPException(404,'Không tìm thấy yêu cầu')
  require_document_access(s,u,d)
  if d.revision!=x.revision:raise HTTPException(409,'Yêu cầu đã thay đổi, vui lòng tải lại')
  if x.decision=='DELETE':
   if d.owner_id!=u.id or d.status!='RESERVATION_REJECTED':raise HTTPException(403,'Chỉ người xin số được xoá yêu cầu bị từ chối')
   s.add(DeletedRecord(document_id=d.id,previous_status=d.status,deleted_by=u.id,reason=x.note.strip() or 'Xoá yêu cầu xin số bị từ chối'))
   d.status='DELETED'
  else:
   if u.role!='CLERK':raise HTTPException(403,'Chỉ Văn thư được duyệt hoặc từ chối xin số')
   if d.status!='RESERVATION_PENDING':raise HTTPException(409,'Yêu cầu không còn chờ duyệt')
   if x.decision=='REJECT':
    if not x.note.strip():raise HTTPException(422,'Nhập lý do từ chối')
    d.status='RESERVATION_REJECTED';d.summary+='\nLý do từ chối: '+x.note.strip()
   else:
    year=d.archive_year or date.today().year
    category='CÔNG VĂN ĐI' if d.direction=='OUT' and d.doc_type=='CÔNG VĂN' else d.doc_type
    d.number=allocate_archive_number(s,year,category);d.symbol=number_symbol(d.doc_type,d.number,date(year,1,1),d.direction)
    d.status='RESERVED_OUTGOING' if d.direction=='OUT' else 'RESERVED_INTERNAL';d.reserved=True
  refresh_search(d);log(s,d,u,'RESERVATION_'+x.decision,x.note.strip());s.commit();return out(d)
 except Exception:
  s.rollback();raise

@app.post('/api/documents/{did}/edit-outgoing')
async def edit_outgoing(did:int,title:str=Form(...),summary:str=Form(''),issuing_agency:str=Form(...),issued_date:Optional[str]=Form(None),file_name:Optional[str]=Form(None),file:Optional[UploadFile]=File(None),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or d.direction!='OUT':raise HTTPException(404,'Không tìm thấy công văn đi')
 require_permission(s,u,'EDIT','Bạn không có quyền sửa công văn đi')
 if not view(u,d,s):raise HTTPException(404,'Không tìm thấy công văn đi')
 if d.status not in {'DRAFT','RETURNED'}:raise HTTPException(409,'Chỉ được sửa bản nháp hoặc văn bản bị trả lại')
 if not title.strip() or not issuing_agency.strip():raise HTTPException(422,'Thiếu thông tin bắt buộc')
 if d.doc_type in {*DOCUMENT_TYPE_CODES,'CÔNG VĂN NỘI BỘ'} and issuing_agency.strip() not in INTERNAL_UNITS:raise HTTPException(422,'Chọn đơn vị tiếp nhận trong trường')
 if not d.file_path and (not file or not file.filename):raise HTTPException(422,'File PDF, DOC hoặc DOCX là bắt buộc')
 d.title=title.strip();d.summary=summary;d.issuing_agency=issuing_agency.strip();d.issued_date=date.fromisoformat(issued_date) if issued_date else d.issued_date;d.file_name=document_file_name(d.title,d.symbol)
 if file and file.filename:
  p=STORE/f'{uuid.uuid4()}.pdf';p.write_bytes(await read_document_upload(file));d.file_path=str(p);d.file_name=document_file_name(d.title,d.symbol)
 source=s.query(DocumentSource).filter_by(document_id=d.id).first()
 if source:source.file_path=d.file_path;source.file_name=d.file_name
 refresh_search(d);log(s,d,u,'EDIT_OUTGOING',f'Cập nhật phiên bản {d.revision}');s.commit();return out(d)
@app.post('/api/documents/{did}/attach-reserved')
async def attach_reserved(did:int,file:UploadFile=File(...),title:Optional[str]=Form(None),issuing_agency:Optional[str]=Form(None),issued_date:Optional[str]=Form(None),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or d.status not in {'RESERVED_INCOMING','RESERVED_OUTGOING','RESERVED_INTERNAL'}:raise HTTPException(404,'Không tìm thấy số đã giữ')
 require_document_access(s,u,d)
 require_permission(s,u,'EDIT','Bạn không có quyền bổ sung văn bản')
 if d.direction in {'OUT','INTERNAL'} and (not issued_date or not (issuing_agency or d.issuing_agency).strip()):raise HTTPException(422,'Ngày phát hành và đơn vị là bắt buộc')
 content=await read_document_upload(file)
 p=STORE/f'{uuid.uuid4()}.pdf';p.write_bytes(content);d.file_path=str(p);d.title=title.strip() if title and title.strip() else d.title;d.file_name=document_file_name(d.title,d.symbol);d.issuing_agency=issuing_agency.strip() if issuing_agency else d.issuing_agency;d.issued_date=date.fromisoformat(issued_date) if issued_date else d.issued_date;d.status='NUMBERED' if d.direction=='IN' else 'DRAFT' if d.direction=='OUT' else 'INTERNAL_NUMBERED';d.reserved=False;refresh_search(d);log(s,d,u,'ATTACH_RESERVED_DOCUMENT',file.filename);s.commit();return out(d)
# Legacy clients must use the reviewed assignment flow as well.
@app.post('/api/documents/{did}/office-opinion')
@app.post('/api/documents/{did}/send-bgh')
@app.post('/api/documents/{did}/bgh-decision')
@app.post('/api/documents/{did}/forward')
def retired_incoming_step(did:int,u=Depends(current)):
 raise HTTPException(409,'Vui lòng tải lại trang và mở Giao việc / phản hồi để thực hiện quy trình duyệt mới')

class ReviewAssignment(BaseModel):
 department_id:int
 responsibility:Literal['Đơn vị chủ trì xử lý','Đơn vị đồng chủ trì xử lý','Đơn vị phối hợp','Đơn vị tiếp nhận thông tin']
class ReviewSubmission(BaseModel):
 purpose:str=Field(default='',max_length=500)
 revision:int=0
 note:str
 assignments:list[ReviewAssignment]=Field(default_factory=list)
 deadline:Optional[date]=None
class ReviewResponse(BaseModel):
 revision:int
 decision:str
 comment:str=''
 purposes:dict[int,Literal['Để báo cáo','Đơn vị chủ trì xử lý','Đơn vị đồng chủ trì xử lý','Đơn vị phối hợp','Đơn vị tiếp nhận thông tin']]=Field(default_factory=dict)
 deadline:Optional[date]=None

def review_document(s,u,did):
 d=s.get(Document,did)
 if not d or d.direction!='IN':raise HTTPException(404,'Không tìm thấy công văn đến')
 require_document_access(s,u,d)
 return d

def latest_review(s,did):
 return s.query(IncomingReview).filter_by(document_id=did).order_by(IncomingReview.revision.desc()).first()

def report_tracking_rows(s:Session,u:User):
 rows=[]
 document_ids=[item[0] for item in s.query(IncomingReview.document_id).distinct()]
 for document_id in document_ids:
  document=s.get(Document,document_id);review=latest_review(s,document_id)
  if not document or not review or not view(u,document,s):continue
  assignments=json.loads(review.assignments);decisions=json.loads(review.decisions)
  deadline=next((item.get('deadline') for item in assignments if item.get('deadline')),None)
  units=[]
  for assignment in assignments:
   decision=decisions.get(str(assignment['department_id']),{})
   units.append({'department_id':assignment['department_id'],'name':assignment['name'],'email':assignment.get('email',''),'status':'RECEIVED' if decision.get('decision')=='REPORT' else 'MISSING','submitted_at':decision.get('at')})
  missing=[item for item in units if item['status']=='MISSING'];received=[item for item in units if item['status']=='RECEIVED']
  days_left=(date.fromisoformat(deadline)-date.today()).days if deadline else None
  rows.append({'document_id':document.id,'symbol':document.symbol,'title':document.title,'revision':review.revision,'deadline':deadline,'days_left':days_left,'warning':bool(missing) and days_left is not None and days_left<=1,'total':len(units),'received_count':len(received),'missing_count':len(missing),'received':received,'missing':missing})
 return sorted(rows,key=lambda item:(item['deadline'] is None,item['deadline'] or '9999-12-31',item['title']))

@app.get('/api/email-templates')
def email_templates(month:Optional[date]=None,unit:str='',u=Depends(current),s:Session=Depends(db)):
 chosen=month or date.today()
 department=s.get(Department,u.department_id) if u.department_id else None
 return {'monthly_reminder':monthly_reminder(chosen),'monthly_report':monthly_report(chosen,unit.strip() or (department.name if department else '[Tên đơn vị]'))}

def incoming_review_email(d,r):
 assignments=json.loads(r.assignments)
 return incoming_email(d.title,[a['name'] for a in assignments],d.symbol,d.issued_date,r.note)

@app.get('/api/reports/tracking')
def report_tracking(u=Depends(current),s:Session=Depends(db)):
 require_office_module(u)
 return report_tracking_rows(s,u)

def review_actor(s,u,r):
 if has_permission(s,u,'BGH_DECIDE'):return 'BGH'
 if u.role in {'DEPARTMENT','DEPARTMENT_HEAD'} and any(a['department_id']==u.department_id for a in json.loads(r.assignments)):return str(u.department_id)
 return None

def review_complete(r):
 decisions=json.loads(r.decisions)
 return all(decisions.get(key,{}).get('decision') in {'APPROVE','REPORT'} for key in ['BGH']+[str(a['department_id']) for a in json.loads(r.assignments)])

@app.get('/api/documents/{did}/review')
def get_review(did:int,u=Depends(current),s:Session=Depends(db)):
 d=review_document(s,u,did);r=latest_review(s,did)
 rounds=s.query(IncomingReview).filter_by(document_id=did).order_by(IncomingReview.revision).all()
 comments=s.query(Action,User).join(User,Action.user_id==User.id).filter(Action.document_id==did,Action.action.like('REVIEW_%')).order_by(Action.id).all()
 mutable=d.status not in {'DELETED','ARCHIVED','FORWARDED'}
 actor=review_actor(s,u,r) if r else None
 return {'email_template':incoming_review_email(d,r) if r else None,'revision':r.revision if r else 0,'note':r.note if r else '', 'purpose':(json.loads(r.assignments)[0].get('purpose','') if r and json.loads(r.assignments) else ''), 'assignments':json.loads(r.assignments) if r else [],'decisions':json.loads(r.decisions) if r else {},'status':d.status,'can_submit':mutable and d.status not in {'RECEIVED','RESERVED_INCOMING'} and has_permission(s,u,'OFFICE_OPINION') and has_permission(s,u,'SEND_BGH'),'can_comment':mutable and bool(r) and (has_permission(s,u,'OFFICE_OPINION') or bool(actor)), 'can_decide':bool(r) and actor=='BGH' and d.status=='PENDING_BGH', 'can_report':bool(r) and actor not in {None,'BGH'} and d.status in {'DEPARTMENT_REVIEW','REVIEW_READY'}, 'can_email':bool(r) and bool(json.loads(r.assignments)) and d.status=='REVIEW_READY' and review_complete(r) and has_permission(s,u,'SEND_BGH'), 'rounds':[{'revision':x.revision,'note':x.note,'assignments':json.loads(x.assignments),'decisions':json.loads(x.decisions)} for x in rounds], 'comments':[{'id':a.id,'user':user.full_name,'action':a.action,'created_at':a.created_at,**{key:value for key,value in json.loads(a.comment).items() if key!='report_file_path'}} for a,user in comments]}

def lock_review(s):
 # Serialize read/modify/write of approvals and revisions, including mail release.
 s.connection().exec_driver_sql('BEGIN IMMEDIATE')

@app.post('/api/documents/{did}/review/submit')
def submit_review(did:int,x:ReviewSubmission,u=Depends(current),s:Session=Depends(db)):
 lock_review(s);d=review_document(s,u,did)
 require_permission(s,u,'OFFICE_OPINION');require_permission(s,u,'SEND_BGH')
 if d.status in {'RECEIVED','RESERVED_INCOMING','FORWARDED','ARCHIVED','DELETED'}:raise HTTPException(409,'Văn bản không ở trạng thái giao việc hoặc sửa phiếu')
 old=latest_review(s,did)
 if x.revision!=(old.revision if old else 0):raise HTTPException(409,'Phiếu đã thay đổi, vui lòng tải lại')
 if not x.note.strip():raise HTTPException(422,'Phải nhập nội dung giao việc')
 assignments=[];seen=set()
 for a in x.assignments:
  dep=s.get(Department,a.department_id)
  if not dep or dep.code=='BGH' or dep.id in seen:raise HTTPException(422,'Đơn vị giao việc không hợp lệ hoặc bị trùng')
  seen.add(dep.id);assignments.append({'department_id':dep.id,'name':dep.name,'email':dep.email,'purpose':x.purpose.strip(),'responsibility':a.responsibility,'deadline':x.deadline.isoformat() if x.deadline else None})
 source=incoming_source(s,d)
 if not source:raise HTTPException(409,'Không tìm thấy bản PDF gốc')
 path=resolve_source_file(source,d)
 if not path:raise HTTPException(404,'Không tìm thấy PDF gốc')
 revision=(old.revision if old else 0)+1
 note=x.note.strip()
 # The attached opinion sheet contains only the office head's written opinion.
 # Assignment roles, recipients, deadline and revision remain structured system data.
 pdf=append_opinion(path,note,u.full_name,d.symbol or '',d.title)
 r=IncomingReview(document_id=did,revision=revision,note=note,assignments=json.dumps(assignments,ensure_ascii=False),decisions='{}',pdf_path=str(pdf),created_by=u.id);s.add(r)
 d.file_path=str(pdf);d.assignee_ids=json.dumps(sorted(seen));d.status='PENDING_BGH'
 log(s,d,u,'REVIEW_SUBMIT',json.dumps({'revision':revision,'comment':note},ensure_ascii=False));s.commit();return {'status':d.status,'revision':revision}

@app.post('/api/documents/{did}/review/respond')
def respond_review(did:int,x:ReviewResponse,u=Depends(current),s:Session=Depends(db)):
 lock_review(s);d=review_document(s,u,did);r=latest_review(s,did)
 if not r or r.revision!=x.revision:raise HTTPException(409,'Vòng duyệt đã thay đổi, vui lòng tải lại')
 if d.status in {'FORWARDED','ARCHIVED','DELETED'}:raise HTTPException(409,'Hồ sơ đã kết thúc phản hồi')
 actor=review_actor(s,u,r)
 if x.decision=='COMMENT':
  if not actor and not has_permission(s,u,'OFFICE_OPINION'):raise HTTPException(403,'Bạn không tham gia xử lý hồ sơ')
  if not x.comment.strip():raise HTTPException(422,'Nhập nội dung bình luận')
 else:
  if x.decision not in {'APPROVE','REJECT'}:raise HTTPException(422,'Quyết định không hợp lệ')
  if actor!='BGH':raise HTTPException(403,'Đơn vị gửi báo cáo thay cho thao tác phê duyệt')
  expected={'PENDING_BGH'}
  if d.status not in expected:raise HTTPException(409,'Chưa đến lượt duyệt hoặc hồ sơ đã được yêu cầu chỉnh sửa')
  if x.decision=='REJECT' and not x.comment.strip():raise HTTPException(422,'Vui lòng nêu nội dung cần chỉnh sửa')
  decisions=json.loads(r.decisions);decisions[actor]={'decision':x.decision,'user':u.full_name,'user_id':u.id,'comment':x.comment.strip(),'at':datetime.utcnow().isoformat()};r.decisions=json.dumps(decisions,ensure_ascii=False)
  if x.decision=='REJECT':d.status='REVIEW_CHANGES'
  elif actor=='BGH':d.status='REVIEW_READY' if review_complete(r) else 'DEPARTMENT_REVIEW'
  elif review_complete(r):d.status='REVIEW_READY'
 d.updated_at=datetime.utcnow()
 log(s,d,u,'REVIEW_'+x.decision,json.dumps({'revision':r.revision,'comment':x.comment.strip(),'actor':actor},ensure_ascii=False));s.commit();return {'status':d.status}

@app.post('/api/documents/{did}/review/report')
async def submit_review_report(did:int,revision:int=Form(...),comment:str=Form(''),file:Optional[UploadFile]=File(None),u=Depends(current),s:Session=Depends(db)):
 lock_review(s);d=review_document(s,u,did);r=latest_review(s,did)
 if not r or revision!=r.revision:raise HTTPException(409,'Phiếu đã thay đổi, vui lòng tải lại')
 actor=review_actor(s,u,r)
 if not actor or actor=='BGH':raise HTTPException(403,'Chỉ đơn vị được phân công mới được gửi báo cáo')
 if d.status not in {'DEPARTMENT_REVIEW','REVIEW_READY'}:raise HTTPException(409,'Chưa đến lượt đơn vị gửi báo cáo')
 if not comment.strip() and (not file or not file.filename):raise HTTPException(422,'Nhập nội dung hoặc đính kèm tệp báo cáo')
 payload={'revision':r.revision,'comment':comment.strip(),'actor':actor}
 if file and file.filename:
  suffix=Path(file.filename).suffix.lower()
  if suffix not in {'.pdf','.doc','.docx'}:raise HTTPException(400,'Báo cáo chỉ nhận tệp PDF, DOC hoặc DOCX')
  content=await file.read(25*1024*1024+1)
  if len(content)>25*1024*1024:raise HTTPException(413,'Tệp báo cáo vượt quá 25 MB')
  assignments=json.loads(r.assignments);deadline=next((item.get('deadline') for item in assignments if item.get('deadline')),None);year=date.fromisoformat(deadline).year if deadline else date.today().year
  folder=ARCHIVE_ROOT/str(year)/'BÁO CÁO'/f'{d.id:06d}';folder.mkdir(parents=True,exist_ok=True)
  safe=''.join(char for char in Path(file.filename).name if char not in '/\\:')
  target=folder/f'{actor}-{uuid.uuid4().hex[:8]}-{safe}';target.write_bytes(content)
  payload.update({'report_file_name':Path(file.filename).name,'report_file_path':str(target)})
 decisions=json.loads(r.decisions);decisions[actor]={'decision':'REPORT','user':u.full_name,'user_id':u.id,'comment':comment.strip(),'at':datetime.utcnow().isoformat()};r.decisions=json.dumps(decisions,ensure_ascii=False)
 if review_complete(r):d.status='REVIEW_READY'
 log(s,d,u,'REVIEW_REPORT',json.dumps(payload,ensure_ascii=False));s.commit();return {'status':d.status}

@app.get('/api/documents/{did}/review/reports/{action_id}')
def download_review_report(did:int,action_id:int,u=Depends(current),s:Session=Depends(db)):
 d=review_document(s,u,did);action=s.get(Action,action_id)
 if not action or action.document_id!=d.id or action.action!='REVIEW_REPORT':raise HTTPException(404,'Không tìm thấy báo cáo')
 payload=json.loads(action.comment);path=Path(payload.get('report_file_path',''))
 valid=False
 if path.is_file():
  if path.parent.resolve()==STORE.resolve():valid=True
  else:
   try:parts=path.resolve().relative_to(ARCHIVE_ROOT.resolve()).parts;valid=len(parts)>=3 and parts[1]=='BÁO CÁO'
   except ValueError:pass
 if not valid:raise HTTPException(404,'Tệp báo cáo không còn tồn tại')
 return FileResponse(path,filename=payload.get('report_file_name') or path.name)

@app.post('/api/documents/{did}/review/email')
def email_review(did:int,x:ReviewResponse,u=Depends(current),s:Session=Depends(db)):
 lock_review(s);d=review_document(s,u,did);r=latest_review(s,did);require_permission(s,u,'SEND_BGH')
 if not r or r.revision!=x.revision or d.status!='REVIEW_READY' or not review_complete(r):raise HTTPException(409,'Chỉ gửi mail khi BGH và tất cả đơn vị đồng ý cùng phiên bản phiếu')
 assignments=json.loads(r.assignments)
 if set(x.purposes)-{a['department_id'] for a in assignments}:raise HTTPException(422,'Mục đích chỉ áp dụng cho đơn vị được giao việc')
 for a in assignments:
  dep=s.get(Department,a['department_id'])
  if not dep or not re.fullmatch(r'[^\s@,;<>]+@[^\s@,;<>]+\.[^\s@,;<>]+',dep.email or ''):raise HTTPException(422,f"Chưa cấu hình email hợp lệ cho {a['name']} trong Cấu hình hệ thống")
  a['email']=dep.email
 recipients=sorted({a['email'] for a in assignments})
 delivery=[{**a,'purpose':a.get('purpose','')} for a in assignments]
 deadline=assignments[0].get('deadline') if assignments else None
 if (x.purposes and any(x.purposes.get(a['department_id'],a.get('purpose',''))!=a.get('purpose','') for a in assignments)) or (x.deadline and x.deadline.isoformat()!=deadline):raise HTTPException(409,'Mục đích hoặc hạn xử lý thay đổi: cần sửa phiếu và trình duyệt lại')
 content=template_html(incoming_review_email(d,r))
 setting=s.query(EmailSetting).first()
 return gmail_draft(d,recipients,d.title,content,setting.signature_html if setting else '')

class GmailConfirmation(BaseModel):
 revision:int
 review_revision:Optional[int]=None

@app.post('/api/documents/{did}/gmail-confirm')
def confirm_gmail(did:int,x:GmailConfirmation,u=Depends(current),s:Session=Depends(db)):
 lock_review(s)
 d=s.get(Document,did)
 if not d:raise HTTPException(404,'Không tìm thấy văn bản')
 require_document_access(s,u,d)
 require_permission(s,u,'SEND_BGH' if d.direction=='IN' else 'EMAIL')
 if d.revision!=x.revision:raise HTTPException(409,'Văn bản đã thay đổi, hãy kiểm tra lại')
 if d.direction=='IN':
  r=latest_review(s,did)
  if not r or r.revision!=x.review_revision or d.status!='REVIEW_READY' or not review_complete(r):raise HTTPException(409,'Phiếu đã thay đổi hoặc chưa đủ điều kiện')
  d.status='FORWARDED'
 elif d.direction=='OUT' and d.status=='SEALED' and d.sealed:d.status='EMAILED'
 elif d.direction=='INTERNAL' and d.status=='INTERNAL_SEALED' and d.sealed:d.status='INTERNAL_PUBLISHED'
 else:raise HTTPException(409,'Văn bản không ở bước chờ gửi email')
 log(s,d,u,'GMAIL_SEND_CONFIRMED','Người dùng xác nhận đã tự gửi thư trong Gmail; hệ thống không xác minh giao thư')
 s.commit();return out(d)

@app.post('/api/sequences/reserve')
def reserve(x:Reserve,u=Depends(current),s:Session=Depends(db)):
 department=reservation_department(s,u,x.department_id)
 if x.doc_type not in DOC_TYPE_CODES:raise HTTPException(422,'Loại văn bản không hợp lệ')
 result=[]
 for _ in range(min(max(x.quantity,1),50)):
  d=Document(direction='OUT',doc_type=x.doc_type,title='Yêu cầu xin số trước',summary='Lý do xin số trước: Xin số theo lô',archive_year=x.year,status='RESERVATION_PENDING',reserved=True,owner_id=u.id,department_id=department.id,issuing_agency=department.name);s.add(d);s.flush();log(s,d,u,'RESERVATION_SUBMIT');result.append(out(d))
 s.commit();return result
@app.post('/api/documents/{did}/action')
def action(did:int,x:Act,u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or not view(u,d,s):raise HTTPException(404,'Không tìm thấy văn bản')
 if x.action=='DELETE':
  require_permission(s,u,'DELETE','Bạn không có quyền xóa văn bản')
  if not x.comment.strip():raise HTTPException(422,'Lý do xóa là bắt buộc')
  if d.status=='DELETED':raise HTTPException(409,'Văn bản đã được xóa')
  s.add(DeletedRecord(document_id=d.id,previous_status=d.status,deleted_by=u.id,reason=x.comment.strip()));d.status='DELETED';log(s,d,u,'SOFT_DELETE',x.comment.strip());s.commit();return out(d)
 if x.action=='ARCHIVE':
  if d.status=='ARCHIVED':raise HTTPException(409,'Văn bản đã được lưu trữ')
  require_permission(s,u,'ARCHIVE','Bạn không có quyền lưu trữ văn bản')
  if d.direction=='OUT' and (d.status not in {'SEALED','EMAILED'} or not d.sealed):raise HTTPException(409,'Văn bản đi phải được đóng dấu trước khi lưu trữ')
  if d.direction=='IN':
   r=latest_review(s,d.id)
   if d.status not in {'REVIEW_READY','FORWARDED'} or not r or not review_complete(r):raise HTTPException(409,'Công văn đến phải được BGH và các đơn vị đồng ý trước khi lưu trữ')
  if d.direction=='INTERNAL' and (d.status not in {'INTERNAL_SEALED','INTERNAL_PUBLISHED'} or not d.sealed):raise HTTPException(409,'Văn bản nội bộ phải được đóng dấu trước khi lưu trữ')
  if not d.file_path:raise HTTPException(400,'Văn bản chưa có PDF để OCR')
  source=incoming_source(s,d) if d.direction=='IN' else None
  archive_path=Path(source.file_path) if source else Path(d.file_path)
  d.ocr_text=ocr_pdf(archive_path)
  archive_file(d,source)
  archive_note=('Lưu không gửi mail' if d.status in {'SEALED','INTERNAL_SEALED','REVIEW_READY'} else 'Lưu sau khi gửi mail')
  d.status='ARCHIVED';refresh_search(d);log(s,d,u,'ARCHIVE',archive_note+(' — '+x.comment if x.comment else ''));s.commit();return out(d)
 raise HTTPException(400,'Thao tác không hợp lệ hoặc đã được thay bằng quy trình chuyên biệt')
@app.post('/api/documents/{did}/restore')
def restore_document(did:int,u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'RESTORE','Bạn không có quyền khôi phục văn bản')
 d=s.get(Document,did);deleted=s.query(DeletedRecord).filter_by(document_id=did).first()
 if not d or d.status!='DELETED' or not deleted:raise HTTPException(404,'Không tìm thấy văn bản đã xóa')
 d.status=deleted.previous_status;log(s,d,u,'RESTORE',deleted.reason);s.delete(deleted);s.commit();return out(d)
@app.get('/api/documents/{did}/history')
def history(did:int,u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or not view(u,d,s):raise HTTPException(404)
 return [{'action':a.action,'comment':a.comment,'user_id':a.user_id,'created_at':a.created_at} for a in s.query(Action).filter_by(document_id=did).order_by(Action.created_at.desc())]
@app.get('/api/documents/{did}/file')
def download(did:int,original:bool=False,u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or not view(u,d,s) or not d.file_path:raise HTTPException(404)
 source=incoming_source(s,d) if original and d.direction=='IN' else None
 path=resolve_source_file(source,d) if source else resolve_managed_file(d)
 # Old databases can contain absolute paths from a previous machine. Repair
 # them lazily once the corresponding archived file is found in this project.
 if source and path and str(path)!=source.file_path:
  source.file_path=str(path);s.commit()
 if not path or not path.exists():raise HTTPException(404,'Không tìm thấy file văn bản')
 if not source and str(path)!=d.file_path:d.file_path=str(path);s.commit()
 return FileResponse(path,filename=d.file_name,headers={'Cache-Control':'no-store'})
@app.post('/api/documents/{did}/rename-file')
def rename_file(did:int,x:RenameFile,u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or not view(u,d,s) or not d.file_path:raise HTTPException(404,'Không tìm thấy file văn bản')
 require_permission(s,u,'RENAME','Bạn không có quyền đổi tên file')
 name=x.name.strip()
 if not name or any(c in name for c in '/\\\0'):raise HTTPException(422,'Tên file không hợp lệ')
 source=Path(d.file_path);suffix=source.suffix or Path(d.file_name or '').suffix
 if suffix and not name.lower().endswith(suffix.lower()):name+=suffix
 target=source.with_name(name)
 if target.exists() and target.resolve()!=source.resolve():raise HTTPException(409,'Tên file đã tồn tại')
 old=d.file_name or source.name;source.rename(target);d.file_path=str(target);d.file_name=name;log(s,d,u,'RENAME_FILE',f'{old} -> {name}');s.commit();return out(d)
@app.get('/api/documents/{did}/pages/{page_number}/preview')
def preview_page(did:int,page_number:int,source:bool=False,u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or not view(u,d,s) or not d.file_path:raise HTTPException(404,'Không tìm thấy PDF')
 source_row=s.query(DocumentSource).filter_by(document_id=did).first() if source else None;pdf_path=source_row.file_path if source_row else d.file_path;doc=pymupdf.open(pdf_path)
 if page_number<1 or page_number>len(doc):doc.close();raise HTTPException(404,'Trang PDF không tồn tại')
 pix=doc[page_number-1].get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False);content=pix.tobytes('png');doc.close();return Response(content=content,media_type='image/png')
@app.get('/api/documents/{did}/preview-info')
def preview_info(did:int,source:bool=False,u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or not view(u,d,s) or not d.file_path:raise HTTPException(404,'Không tìm thấy PDF')
 source_row=s.query(DocumentSource).filter_by(document_id=did).first() if source else None;pdf_path=source_row.file_path if source_row else d.file_path;doc=pymupdf.open(pdf_path);pages=[{'page':i+1,'width':round(p.rect.width,2),'height':round(p.rect.height,2)} for i,p in enumerate(doc)];doc.close();return {'page_count':len(pages),'pages':pages}
@app.post('/api/documents/{did}/digital-sign')
async def digital_sign(did:int,page:int=Form(1),x_percent:float=Form(...),y_percent:float=Form(...),note:str=Form(''),file:Optional[UploadFile]=File(None),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or not view(u,d,s):raise HTTPException(404,'Không tìm thấy văn bản')
 if d.direction!='OUT':raise HTTPException(409,'Công văn đến không có bước ký số')
 require_permission(s,u,'SIGN','Bạn không có quyền ký số công văn đi')
 bgh=d.status=='PENDING_OUT_BGH_SIGN'
 if bgh:
  if u.role not in {'ADMIN','BGH'}:raise HTTPException(403,'Chỉ BGH được ký bước này')
  if not note.strip():raise HTTPException(422,'Ý kiến duyệt là bắt buộc trước khi ký số')
 elif d.status in {'DRAFT','RETURNED','PENDING_OFFICE_SIGN'}:
  if u.role not in {'ADMIN','DEPARTMENT','DEPARTMENT_HEAD','OFFICE_HEAD'}:raise HTTPException(403,'Chỉ tài khoản đơn vị được ký bước này')
  if u.role!='ADMIN' and (u.department_id is None or u.department_id!=d.department_id):raise HTTPException(403,'Chỉ được ký văn bản của đơn vị mình')
 else:raise HTTPException(409,'Văn bản không ở bước chờ ký')
 if file:raise HTTPException(409,'Hãy sửa PDF ở bước bản nháp trước khi ký')
 if d.sealed:raise HTTPException(409,'Văn bản đã đóng dấu nên không thể thay đổi chữ ký')
 if not d.file_path or not Path(d.file_path).exists():raise HTTPException(400,'Văn bản chưa có tệp PDF')
 source=s.query(DocumentSource).filter_by(document_id=d.id).first()
 if not source:source=DocumentSource(document_id=d.id,file_path=d.file_path,file_name=d.file_name or 'van-ban.pdf');s.add(source);s.flush()
 if not Path(source.file_path).exists():raise HTTPException(400,'Không tìm thấy PDF nguồn')
 signed_path=render_signature(Path(d.file_path),max(0,min(100,x_percent)),max(0,min(100,y_percent)),max(1,page));digest=hashlib.sha256(signed_path.read_bytes()).hexdigest();d.file_path=str(signed_path)
 sig=DigitalSignature(document_id=d.id,user_id=u.id,page=max(1,page),x_percent=max(0,min(100,x_percent)),y_percent=max(0,min(100,y_percent)),file_hash=digest);s.add(sig)
 d.signed_personal=True;d.signature_hash=digest;d.status='READY_FOR_CLERK' if bgh else 'UNIT_SIGNED'
 if bgh:d.received_date=date.today()
 detail=json.dumps({'decision':'APPROVE_AND_SIGN','revision':d.revision,'note':note.strip(),'page':sig.page,'x_percent':sig.x_percent,'y_percent':sig.y_percent},ensure_ascii=False) if bgh else f'Trang {sig.page}, vị trí {sig.x_percent}% / {sig.y_percent}% (giả lập)'
 log(s,d,u,'BGH_DIGITAL_SIGN' if bgh else 'UNIT_DIGITAL_SIGN',detail);s.commit()
 return {'status':'signed_simulation','signature_id':sig.id,'page':sig.page,'x_percent':sig.x_percent,'y_percent':sig.y_percent,'file_hash':digest}
@app.get('/api/documents/{did}/digital-signature')
def digital_signature_info(did:int,u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or not view(u,d,s):raise HTTPException(404,'Không tìm thấy văn bản')
 sig=s.query(DigitalSignature).filter_by(document_id=did).order_by(DigitalSignature.created_at.desc()).first()
 return {'signed':bool(sig),'page':sig.page if sig else 1,'x_percent':sig.x_percent if sig else 55,'y_percent':sig.y_percent if sig else 68,'locked':d.sealed}
@app.post('/api/documents/{did}/digital-seal')
def digital_seal(did:int,seal_type:str=Form(...),page:int=Form(1),x_percent:float=Form(...),y_percent:float=Form(...),u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or not view(u,d,s):raise HTTPException(404,'Không tìm thấy văn bản')
 if d.direction not in {'OUT','INTERNAL'}:raise HTTPException(409,'Công văn đến không có bước đóng dấu')
 require_permission(s,u,'SEAL','Bạn không có quyền đóng dấu')
 if d.direction=='OUT' and not d.signed_personal:raise HTTPException(409,'Văn bản phải được ký số trước khi đóng dấu')
 expected_status='OUT_NUMBERED' if d.direction=='OUT' else 'INTERNAL_APPROVED'
 if d.status!=expected_status:raise HTTPException(409,'Văn bản đi phải được cấp số; văn bản nội bộ phải được Chánh Văn phòng phê duyệt trước khi đóng dấu')
 if d.sealed:raise HTTPException(409,'Văn bản này đã được đóng dấu')
 if seal_type not in {'SCHOOL','OFFICE'}:raise HTTPException(400,'Loại con dấu không hợp lệ')
 if not d.file_path or not Path(d.file_path).exists():raise HTTPException(400,'Văn bản chưa có tệp PDF')
 sealed_path=render_seal(Path(d.file_path),max(0,min(100,x_percent)),max(0,min(100,y_percent)),max(1,page),seal_type);digest=hashlib.sha256(sealed_path.read_bytes()).hexdigest();d.file_path=str(sealed_path);seal=DigitalSeal(document_id=d.id,user_id=u.id,seal_type=seal_type,page=max(1,page),x_percent=max(0,min(100,x_percent)),y_percent=max(0,min(100,y_percent)),file_hash=digest)
 s.add(seal);d.sealed=True;d.status='SEALED' if d.direction=='OUT' else 'INTERNAL_SEALED';log(s,d,u,'DIGITAL_SEAL',f'{seal_type}, trang {seal.page}, vị trí {seal.x_percent}% / {seal.y_percent}% (giả lập)');s.commit()
 return {'status':'sealed_simulation','seal_id':seal.id,'seal_type':seal.seal_type,'page':seal.page,'x_percent':seal.x_percent,'y_percent':seal.y_percent,'file_hash':digest}
@app.get('/api/config')
def config(u=Depends(current),s:Session=Depends(db)):
 effective=[code for code in PERMISSION_LABELS if has_permission(s,u,code)]
 rank={code:i for i,(code,_) in enumerate(SCHOOL_DEPARTMENTS)};departments=[d for d in s.query(Department).all() if d.code in rank or d.code=='BGH'];departments.sort(key=lambda d:(d.code=='BGH',rank.get(d.code,999),d.name))
 if u.role!='ADMIN':return {'departments':[{'id':d.id,'name':d.name,'code':d.code,'email':d.email} for d in departments],'sequences':[],'users':[],'roles':[],'email':{},'permissions':effective}
 email=s.query(EmailSetting).first();roles=[
  {'code':'ADMIN','name':'Quản trị hệ thống','permissions':['Quản lý người dùng','Cấu hình hệ thống','Xem toàn bộ hồ sơ','Xóa và khôi phục']},
  {'code':'BGH','name':'Ban Giám hiệu','permissions':['Xem hồ sơ được trình','Duyệt hoặc từ chối','Theo dõi xử lý']},
  {'code':'OFFICE_HEAD','name':'Chánh Văn phòng','permissions':['Cho ý kiến công văn đến','Gửi Ban Giám hiệu','Phê duyệt văn bản nội bộ','Ký hoặc trả công văn đi']},
  {'code':'CLERK','name':'Văn thư','permissions':['Tiếp nhận và tạo văn bản','Sửa và cấp số','Chuyển đơn vị xử lý','Đóng dấu','Phát hành qua email','Lưu trữ','Xóa có lý do']},
  {'code':'DEPARTMENT_HEAD','name':'Trưởng đơn vị','permissions':['Tạo văn bản','Ký số đơn vị','Trình BGH']},
  {'code':'DEPARTMENT','name':'Đơn vị xử lý','permissions':['Xem công văn đến được phân công','Tạo, sửa và ký công văn đi của đơn vị','Trình BGH','Thảo luận và gửi báo cáo']},
 ]
 overrides={(x.role,x.permission):x.enabled for x in s.query(RolePermission)}
 for role in roles:
  role['permission_options']=[{'code':code,'name':PERMISSION_LABELS[code],'enabled':overrides.get((role['code'],code),True)} for code in ROLE_DEFAULTS.get(role['code'],set())]
  role['permissions']=[x['name'] for x in role['permission_options'] if x['enabled']]
 return {'departments':[{'id':d.id,'name':d.name,'code':d.code,'email':d.email} for d in departments],'sequences':[{'year':x.year,'doc_type':x.doc_type,'current':x.current,'prefix':x.prefix} for x in s.query(Sequence)],'users':[{'id':x.id,'username':x.username,'full_name':x.full_name,'role':x.role,'department_id':x.department_id,'active':x.active} for x in s.query(User)] if has_permission(s,u,'MANAGE_USERS') else [],'roles':roles,'email':({'smtp_host':email.smtp_host,'smtp_port':email.smtp_port,'username':email.username,'sender_name':email.sender_name,'sender_email':email.sender_email,'signature_html':email.signature_html,'reminder_subject':email.reminder_subject,'reminder_html':email.reminder_html,'use_tls':email.use_tls,'enabled':email.enabled} if email else {'smtp_host':'','smtp_port':587,'username':'','sender_name':'Đại học Hùng Vương','sender_email':'','signature_html':'','reminder_subject':'Nhắc hạn nộp báo cáo: {title}','reminder_html':'<p>Đơn vị chưa nộp báo cáo cho hồ sơ <b>{title}</b>.</p><p>Hạn nộp: <b>{deadline}</b>.</p>','use_tls':True,'enabled':False}),'permissions':effective}
class DepartmentEmailConfig(BaseModel):
 email:str=''
@app.put('/api/config/departments/{did}/email')
def configure_department_email(did:int,x:DepartmentEmailConfig,u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'MANAGE_EMAIL')
 dep=s.get(Department,did)
 if not dep:raise HTTPException(404,'Không tìm thấy đơn vị')
 value=x.email.strip()
 if value and not re.fullmatch(r'[^\s@,;<>]+@[^\s@,;<>]+\.[^\s@,;<>]+',value):raise HTTPException(422,'Email không hợp lệ')
 dep.email=value;s.commit();return {'id':dep.id,'email':dep.email}

class CreateUnitAccount(BaseModel):
 username:str=Field(min_length=1,max_length=100)
 full_name:str=Field(min_length=1,max_length=200)
 password:str=Field(min_length=8,max_length=128)
 role:Literal['DEPARTMENT','DEPARTMENT_HEAD']='DEPARTMENT'
 department_id:int
@app.post('/api/config/users')
def create_unit_account(x:CreateUnitAccount,u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'MANAGE_USERS')
 if not x.username.strip() or not x.full_name.strip():raise HTTPException(422,'Tên tài khoản và họ tên là bắt buộc')
 dep=s.get(Department,x.department_id)
 if not dep or dep.code not in {code for code,_ in SCHOOL_DEPARTMENTS}:raise HTTPException(422,'Đơn vị không thuộc danh mục đơn vị của trường')
 if s.query(User).filter_by(username=x.username.strip()).first():raise HTTPException(409,'Tên tài khoản đã tồn tại')
 user=User(username=x.username.strip(),full_name=x.full_name.strip(),password_hash=pwd.hash(x.password),role=x.role,department_id=x.department_id,active=True)
 s.add(user);s.commit();return {'id':user.id,'username':user.username,'role':user.role}
@app.put('/api/config/users/{uid}')
def update_user_config(uid:int,x:UserConfig,u=Depends(current),s:Session=Depends(db)):
 if u.role!='ADMIN':raise HTTPException(403,'Chỉ Admin được cấu hình người dùng')
 target=s.get(User,uid)
 if not target:raise HTTPException(404,'Không tìm thấy người dùng')
 if x.role not in {'ADMIN','BGH','OFFICE_HEAD','CLERK','DEPARTMENT','DEPARTMENT_HEAD'}:raise HTTPException(400,'Vai trò không hợp lệ')
 if target.username=='admin' and (not x.active or x.role!='ADMIN'):raise HTTPException(409,'Tài khoản admin luôn hoạt động và luôn có toàn quyền')
 if target.username!='admin' and x.role=='ADMIN':raise HTTPException(409,'Không thể gán vai trò quản trị tối cao cho tài khoản khác')
 dep=s.get(Department,x.department_id) if x.department_id else None
 school_codes={code for code,_ in SCHOOL_DEPARTMENTS}
 if x.role in {'DEPARTMENT','DEPARTMENT_HEAD'} and (not dep or dep.code not in school_codes):raise HTTPException(422,'Tài khoản đơn vị phải chọn một đơn vị trong danh mục của trường')
 if x.role=='BGH' and (not dep or dep.code!='BGH'):raise HTTPException(422,'Tài khoản Ban Giám hiệu phải thuộc Ban Giám hiệu')
 if x.role in {'OFFICE_HEAD','CLERK'} and (not dep or dep.code!='VP'):raise HTTPException(422,'Tài khoản Văn phòng phải thuộc Văn phòng Trường')
 target.full_name=x.full_name.strip();target.role=x.role;target.department_id=x.department_id;target.active=x.active;s.commit();return {'status':'updated'}
@app.put('/api/config/email')
def update_email_config(x:EmailConfig,u=Depends(current),s:Session=Depends(db)):
 if u.role!='ADMIN':raise HTTPException(403,'Chỉ Admin được cấu hình email')
 if x.smtp_port<1 or x.smtp_port>65535:raise HTTPException(422,'Cổng SMTP không hợp lệ')
 x.signature_html=constrain_signature_images(clean_email_html(x.signature_html))
 x.reminder_html=clean_email_html(x.reminder_html)
 x.reminder_subject=x.reminder_subject.strip()
 if not has_email_content(x.signature_html):raise HTTPException(422,'Chữ ký email là bắt buộc')
 if not x.reminder_subject or not has_email_content(x.reminder_html):raise HTTPException(422,'Tiêu đề và nội dung email nhắc hạn là bắt buộc')
 email=s.query(EmailSetting).first()
 if not email:email=EmailSetting();s.add(email)
 for key,value in x.model_dump().items():setattr(email,key,value)
 s.commit();return {'status':'updated'}
@app.put('/api/config/permissions')
def update_permissions(x:PermissionConfig,u=Depends(current),s:Session=Depends(db)):
 if u.role!='ADMIN':raise HTTPException(403,'Chỉ Admin được cấu hình phân quyền')
 if x.role not in ROLE_DEFAULTS:raise HTTPException(400,'Vai trò không hợp lệ')
 if x.role=='ADMIN':raise HTTPException(409,'Quyền Quản trị hệ thống được khóa để tránh mất quyền quản trị')
 invalid=set(x.permissions)-set(PERMISSION_LABELS)
 if invalid:raise HTTPException(400,f'Quyền không hợp lệ: {", ".join(sorted(invalid))}')
 outside_role=set(x.permissions)-ROLE_DEFAULTS[x.role]
 if outside_role:raise HTTPException(400,'Không thể gán quyền nghiệp vụ ngoài phạm vi của vai trò')
 selected=set(x.permissions)
 for code in PERMISSION_LABELS:
  row=s.query(RolePermission).filter_by(role=x.role,permission=code).first()
  if not row:row=RolePermission(role=x.role,permission=code);s.add(row)
  row.enabled=code in selected
 s.commit();return {'status':'updated','role':x.role,'permissions':sorted(selected)}
@app.get('/api/activity')
def activity(limit:int=200,u=Depends(current),s:Session=Depends(db)):
 if u.role!='ADMIN':raise HTTPException(403,'Chỉ Admin được xem nhật ký toàn hệ thống')
 rows=s.query(Action,User,Document).join(User,Action.user_id==User.id).join(Document,Action.document_id==Document.id).order_by(Action.created_at.desc()).limit(max(1,min(limit,1000))).all()
 return [{'id':a.id,'created_at':a.created_at,'user_id':user.id,'username':user.username,'full_name':user.full_name,'role':user.role,'action':a.action,'comment':a.comment,'document_id':doc.id,'symbol':doc.symbol,'title':doc.title,'direction':doc.direction,'status':doc.status} for a,user,doc in rows]

def send_due_report_reminders(s:Session,u:User):
 if not has_permission(s,u,'EMAIL'):return
 setting=s.query(EmailSetting).first()
 if not setting or not setting.enabled:return
 for item in report_tracking_rows(s,u):
  if item['days_left']!=1 or not item['missing']:continue
  marker=f"{item['revision']}:{item['deadline']}"
  if s.query(Action).filter_by(document_id=item['document_id'],action='REVIEW_DEADLINE_REMINDER',comment=marker).first():continue
  recipients=sorted({unit['email'] for unit in item['missing'] if re.fullmatch(r'[^\s@,;<>]+@[^\s@,;<>]+\.[^\s@,;<>]+',unit['email'] or '')})
  if not recipients:continue
  document=s.get(Document,item['document_id']);values={'title':document.title,'symbol':document.symbol or 'Chưa cấp số','deadline':date.fromisoformat(item['deadline']).strftime('%d/%m/%Y'),'missing_units':', '.join(unit['name'] for unit in item['missing'])}
  subject=setting.reminder_subject
  content=setting.reminder_html
  for key,value in values.items():subject=subject.replace('{'+key+'}',value);content=content.replace('{'+key+'}',html.escape(value))
  try:send_document_email(s,document,recipients,subject,content,setting.signature_html)
  except HTTPException:continue
  log(s,document,u,'REVIEW_DEADLINE_REMINDER',marker);s.commit()

@app.get('/api/notifications')
def notifications(u=Depends(current),s:Session=Depends(db)):
 # Email is composed and sent by the user in Gmail.
 tasks={
  'RECEIVED':('ASSIGN_NUMBER','Công văn đến mới','Cần vào số công văn đến'),
  'NUMBERED':('OFFICE_OPINION','Công văn chờ ý kiến','Chánh Văn phòng cần ghi ý kiến xử lý'),
  'OFFICE_OPINION_COMPLETED':('SEND_BGH','Hồ sơ chờ trình BGH','Chánh Văn phòng gửi hồ sơ đã có ý kiến đến Ban Giám hiệu'),
  'PENDING_BGH':('BGH_DECIDE','Văn bản chờ BGH xử lý','Cần duyệt hoặc từ chối văn bản'),
  'REVIEW_CHANGES':('OFFICE_OPINION','Hồ sơ cần chỉnh sửa','Có phản hồi yêu cầu sửa phiếu giao việc'),
  'REVIEW_READY':('SEND_BGH','Hồ sơ đã đồng thuận','BGH và tất cả đơn vị đã đồng ý, có thể gửi email'),
  'DEPARTMENT_REVIEW':(None,'Cần đọc và phản hồi','BGH đã duyệt, đơn vị cần đồng ý hoặc phản hồi'),
  'BGH_APPROVED':('FORWARD','Văn bản đã được duyệt','Cần chuyển đến đơn vị xử lý'),
  'FORWARDED':(None,'Văn bản được phân công','Đơn vị của bạn được giao xử lý văn bản'),
  'PENDING_OFFICE_SIGN':('SIGN','Công văn đi chờ ký','Chánh Văn phòng cần xem xét và ký số'),
  'UNIT_SIGNED':('EDIT','Đơn vị đã ký','Cần gửi BGH duyệt'),
  'PENDING_OUT_BGH_APPROVAL':('BGH_DECIDE','Công văn chờ BGH duyệt','Văn bản đã được lãnh đạo đơn vị ký; BGH chỉ cần duyệt'),
  'PENDING_OUT_BGH_SIGN':('BGH_DECIDE','Công văn chờ BGH ký','Văn bản chưa có chữ ký; BGH cần duyệt và ký số'),
  'READY_FOR_CLERK':('ASSIGN_NUMBER','Đã hoàn tất ký duyệt','Văn thư cần vào số và ngày'),
  'OUT_NUMBERED':('SEAL','Công văn đi đã cấp số','Cần đóng dấu văn bản'),
  'SEALED':('EMAIL','Văn bản đã đóng dấu','Cần phát hành văn bản qua email'),
  'INTERNAL_NUMBERED':('EDIT','Văn bản nội bộ đã cấp số','Cần trình Chánh Văn phòng phê duyệt'),
  'PENDING_INTERNAL_APPROVAL':('INTERNAL_APPROVE','Văn bản nội bộ chờ duyệt','Chánh Văn phòng cần phê duyệt văn bản'),
  'INTERNAL_APPROVED':('SEAL','Văn bản nội bộ đã được duyệt','Cần đóng dấu văn bản'),
  'INTERNAL_SEALED':('EMAIL','Văn bản nội bộ đã đóng dấu','Cần phát hành văn bản'),
  'RESERVATION_PENDING':('ASSIGN_NUMBER','Xin số chờ duyệt','Văn thư cần duyệt hoặc từ chối yêu cầu xin số'),
  'RESERVATION_REJECTED':('RESERVE_NUMBER','Xin số bị từ chối','Chỉnh sửa và trình lại hoặc xoá yêu cầu'),
  'RETURNED':('EDIT','Văn bản bị trả lại','Cần chỉnh sửa và trình lại'),
 }
 rows=[]
 for d in s.query(Document).filter(Document.status.in_(tasks)).order_by(Document.updated_at.desc()).limit(100):
  permission,title,message=tasks[d.status]
  assigned=u.department_id in json.loads(d.assignee_ids or '[]')
  if d.status=='RESERVATION_PENDING':
   if u.role!='CLERK':continue
  elif d.status=='RESERVATION_REJECTED':
   if d.owner_id!=u.id:continue
  elif d.status=='RETURNED':
   if u.role!='ADMIN' and u.department_id!=d.department_id:continue
  elif d.status in {'FORWARDED','DEPARTMENT_REVIEW'}:
   if u.role!='ADMIN' and not assigned:continue
  elif not has_permission(s,u,permission):continue
  if not view(u,d,s):continue
  rows.append({'id':f'{d.id}:{d.status}:{d.updated_at.isoformat()}','document_id':d.id,'direction':d.direction,'status':d.status,'symbol':d.symbol,'document_title':d.title,'title':title,'message':message,'created_at':d.updated_at})
 for item in report_tracking_rows(s,u):
  if not item['warning']:continue
  own_missing=next((unit for unit in item['missing'] if unit['department_id']==u.department_id),None)
  if own_missing:
   title='Sắp đến hạn nộp báo cáo';message=f"Đơn vị của bạn chưa nộp báo cáo; hạn {item['deadline']}";direction='IN'
  elif has_permission(s,u,'VIEW_ALL'):
   title='Cảnh báo tiến độ báo cáo';message=f"Đã nhận {item['received_count']}/{item['total']}; chưa nhận: "+', '.join(unit['name'] for unit in item['missing']);direction='REPORT'
  else:continue
  rows.append({'id':f"report:{item['document_id']}:{item['revision']}:{item['deadline']}:{u.id}",'document_id':item['document_id'],'direction':direction,'status':'REPORT_DEADLINE_WARNING','symbol':item['symbol'],'document_title':item['title'],'title':title,'message':message,'created_at':datetime.combine(date.fromisoformat(item['deadline']),datetime.min.time())})
 for d in s.query(Document).filter(Document.status.in_({'PENDING_ARCHIVE_REVIEW','ARCHIVE_RETURNED'})).order_by(Document.updated_at.desc()):
  if d.status=='PENDING_ARCHIVE_REVIEW' and has_permission(s,u,'ARCHIVE'):
   rows.append({'id':f'{d.id}:{d.status}:{d.updated_at.isoformat()}','document_id':d.id,'direction':'ARCHIVE','status':d.status,'symbol':d.symbol,'document_title':d.title,'title':'Hồ sơ chờ duyệt lưu trữ','message':f'{d.issuing_agency} đã tải lên {d.file_name}','created_at':d.updated_at})
  elif d.status=='ARCHIVE_RETURNED' and d.owner_id==u.id:
   rows.append({'id':f'{d.id}:{d.status}:{d.updated_at.isoformat()}','document_id':d.id,'direction':'ARCHIVE','status':d.status,'symbol':d.symbol,'document_title':d.title,'title':'Hồ sơ lưu trữ bị trả lại','message':'Vui lòng xem lý do và tải phiên bản mới','created_at':d.updated_at})
 return rows
@app.post('/api/documents/{did}/email')
def email(did:int,recipients:list[str],u=Depends(current),s:Session=Depends(db)):
 d=s.get(Document,did)
 if not d or not view(u,d,s):raise HTTPException(404)
 if d.direction=='IN':raise HTTPException(409,'Công văn đến phải gửi email qua luồng giao việc đã được tất cả đồng ý')
 require_permission(s,u,'EMAIL','Bạn không có quyền gửi email văn bản')
 log(s,d,u,'EMAIL',', '.join(recipients));s.commit();return {'status':'simulated','recipients':recipients}

@app.post('/api/ai/chat')
def ai_chat(x:AIChatRequest,u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'AI_CHAT','Bạn không có quyền dùng trợ lý AI')
 try:return chat_search(message=x.message,history=[m.model_dump() for m in x.history],top_k=5)
 except FileNotFoundError as exc:raise HTTPException(503,str(exc)) from exc
 except Exception as exc:raise HTTPException(500,f'Không thể chạy trợ lý AI: {exc}') from exc
@app.post('/api/ai/search-documents')
def ai_search_documents(x:AIFileSearchRequest,u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'AI_CHAT','Bạn không có quyền tra cứu bằng AI')
 query=x.query.strip()
 if not query:raise HTTPException(422,'Vui lòng nhập nội dung cần tìm')
 try:
  candidates=ai_document_search(query,top_k=10)
  from app.embedder import get_embedding_model
  from app.reranker import get_reranker
  get_embedding_model.cache_clear();get_reranker.cache_clear()
  if hasattr(torch.backends,'mps') and torch.backends.mps.is_available():torch.mps.empty_cache()
  return {'query':query,'results':map_query_to_documents(query,candidates)}
 except FileNotFoundError as exc:raise HTTPException(503,str(exc)) from exc
 except Exception as exc:raise HTTPException(500,f'Không thể tra cứu bằng AI: {exc}') from exc
@app.get('/api/ai/status')
def ai_status(u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'AI_CHAT','Bạn không có quyền dùng trợ lý AI')
 return qwen_status()
@app.get('/api/ai/documents')
def ai_documents(u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'AI_CHAT','Bạn không có quyền dùng trợ lý AI')
 return get_indexed_documents()
@app.post('/api/ai/analyze')
def ai_analyze(x:AIAnalysisRequest,u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'AI_CHAT','Bạn không có quyền dùng trợ lý AI')
 available={item['id'] for item in get_indexed_documents()}
 if not set(x.document_ids)<=available:raise HTTPException(404,'Có văn bản đã chọn không còn trong chỉ mục')
 try:return chat_search(message=x.message,history=[m.model_dump() for m in x.history],document_ids=x.document_ids,top_k=10)
 except FileNotFoundError as exc:raise HTTPException(503,str(exc)) from exc
 except Exception as exc:raise HTTPException(500,f'Không thể phân tích văn bản: {exc}') from exc
@app.post('/api/ai/reports/generate')
def ai_report_generate(x:AIReportRequest,u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'AI_CHAT','Bạn không có quyền lập báo cáo AI')
 indexed={item['id']:item for item in get_indexed_documents()}
 missing=[did for did in x.document_ids if did not in indexed]
 if missing:raise HTTPException(404,'Có văn bản nguồn không còn trong chỉ mục AI')
 chunks=fetch_chunks_for_documents(x.document_ids,limit=240)
 grouped={did:[] for did in x.document_ids}
 for chunk in chunks:grouped.get(chunk['document_id'],[]).append(chunk)
 sources=[];warnings=[];blocks=[]
 for number,did in enumerate(x.document_ids,1):
  item=indexed[did];rows=grouped[did]
  sources.append({'id':did,'file_name':item['file_name'],'category':item['category'],'chunk_count':len(rows),'readable':bool(rows),'source_number':number})
  if not rows:
   warnings.append(f"{item['file_name']}: chưa có nội dung OCR để tổng hợp")
   continue
  parts=[]
  for row in rows:
   page=row.get('page_no') or 'không xác định'
   content=' '.join((row.get('content') or '').split())
   if content:parts.append(f'[VB{number}, trang {page}] {content}')
  blocks.append(f"VĂN BẢN VB{number}: {item['file_name']}\n"+'\n'.join(parts))
 evidence='\n\n'.join(blocks)
 if not evidence:raise HTTPException(422,'Các văn bản đã chọn chưa có nội dung OCR để lập báo cáo')
 # Keep the local model context bounded while preserving the source order and citations.
 evidence=evidence[:60000]
 outline=[line.strip() for line in x.outline if line.strip()] or ['Phạm vi và tài liệu sử dụng','Tổng hợp các nội dung chính','Các yêu cầu, trách nhiệm và thời hạn','Điểm khác biệt hoặc nội dung cần làm rõ','Danh mục văn bản nguồn']
 level='ngắn gọn, ưu tiên ý cần quyết định' if x.detail_level=='SHORT' else 'đầy đủ, có phân nhóm nội dung'
 request=(f"Soạn BẢN THẢO BÁO CÁO TỔNG HỢP phục vụ {x.audience}; mức chi tiết: {level}.\n"
          f"Yêu cầu báo cáo: {x.requirement}\nĐề cương bắt buộc:\n"+'\n'.join(f'{i}. {name}' for i,name in enumerate(outline,1))+
          "\nMỗi kết luận phải dẫn [VBn, trang x]. Không chép nguyên khối OCR. "
          "Không suy diễn văn bản đã thực hiện, chưa hoàn thành hay quá hạn vì dữ liệu chỉ phản ánh nội dung tài liệu.")
 if x.instruction.strip():request+=f'\nYêu cầu chỉnh sửa bản nháp: {x.instruction.strip()}'
 try:draft=generate_grounded_report(request,evidence,[m.model_dump() for m in x.history])
 except FileNotFoundError as exc:raise HTTPException(503,str(exc)) from exc
 except Exception as exc:raise HTTPException(500,f'Không thể lập báo cáo tổng hợp: {exc}') from exc
 return {'draft':draft,'sources':sources,'warnings':warnings,'document_count':len(sources)}
def markdown_inline_parts(value:str):
 """Split the small Markdown subset used by generated reports."""
 parts=[];position=0
 for match in re.finditer(r'(\*\*([^*]+)\*\*|\*([^*]+)\*)',value):
  if match.start()>position:parts.append((value[position:match.start()],False,False))
  parts.append((match.group(2) or match.group(3),bool(match.group(2)),bool(match.group(3))))
  position=match.end()
 if position<len(value):parts.append((value[position:],False,False))
 return parts or [(value,False,False)]
def add_markdown_runs(paragraph,value:str):
 for text,bold,italic in markdown_inline_parts(value):
  run=paragraph.add_run(text);run.bold=bold;run.italic=italic
def markdown_to_docx(content:str,doc:WordDocument):
 for raw in content.splitlines():
  line=raw.strip()
  if not line:doc.add_paragraph();continue
  if re.fullmatch(r'(?:---+|___+|\*\*\*+)',line):doc.add_paragraph();continue
  heading=re.match(r'^(#{1,4})\s+(.+)$',line)
  bullet=re.match(r'^[-•]\s+(.+)$',line)
  numbered=re.match(r'^(\d+[.)])\s+(.+)$',line)
  if heading:
   paragraph=doc.add_heading(level=min(len(heading.group(1)),3));add_markdown_runs(paragraph,heading.group(2))
  elif bullet:
   paragraph=doc.add_paragraph(style='List Bullet');add_markdown_runs(paragraph,bullet.group(1))
  elif numbered:
   paragraph=doc.add_paragraph(style='List Number');add_markdown_runs(paragraph,numbered.group(2))
  else:
   paragraph=doc.add_paragraph();add_markdown_runs(paragraph,line)
def markdown_inline_html(value:str)->str:
 escaped=html.escape(value)
 escaped=re.sub(r'\*\*([^*]+)\*\*',r'<strong>\1</strong>',escaped)
 return re.sub(r'\*([^*]+)\*',r'<em>\1</em>',escaped)
def markdown_to_html(content:str)->str:
 blocks=[];in_list=False
 def close_list():
  nonlocal in_list
  if in_list:blocks.append('</ul>');in_list=False
 for raw in content.splitlines():
  line=raw.strip()
  if not line:close_list();continue
  if re.fullmatch(r'(?:---+|___+|\*\*\*+)',line):close_list();blocks.append('<hr>');continue
  heading=re.match(r'^(#{1,4})\s+(.+)$',line)
  bullet=re.match(r'^[-•]\s+(.+)$',line)
  if heading:
   close_list();level=min(len(heading.group(1))+1,4);blocks.append(f'<h{level}>{markdown_inline_html(heading.group(2))}</h{level}>')
  elif bullet:
   if not in_list:blocks.append('<ul>');in_list=True
   blocks.append(f'<li>{markdown_inline_html(bullet.group(1))}</li>')
  else:
   close_list();blocks.append(f'<p>{markdown_inline_html(line)}</p>')
 close_list();return ''.join(blocks)
@app.post('/api/ai/reports/export/{file_type}')
def ai_report_export(file_type:str,x:AIReportExport,u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'AI_CHAT','Bạn không có quyền xuất báo cáo AI')
 if file_type=='docx':
  doc=WordDocument();doc.add_heading(x.title,0)
  markdown_to_docx(x.content,doc)
  output=io.BytesIO();doc.save(output);return Response(output.getvalue(),media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':'attachment; filename="bao-cao-dhv.docx"'})
 if file_type=='pdf':
  pdf=pymupdf.open();lines=x.content.splitlines();parts=[lines[i:i+38] for i in range(0,len(lines),38)] or [[]]
  css='body{font-family:sans-serif;font-size:11pt;line-height:1.45;color:#202c3d}h2,h3,h4{color:#193b76;margin:12px 0 6px}p{margin:0 0 7px}li{margin-bottom:4px}hr{border:0;border-top:1px solid #ccd3dc;margin:12px 0}'
  for index,part in enumerate(parts):
   page=pdf.new_page();heading=f'<h1>{html.escape(x.title)}</h1>' if index==0 else '';page.insert_htmlbox(pymupdf.Rect(50,45,545,797),heading+markdown_to_html('\n'.join(part)),css=css,scale_low=.75)
  data=pdf.tobytes();pdf.close();return Response(data,media_type='application/pdf',headers={'Content-Disposition':'attachment; filename="bao-cao-dhv.pdf"'})
 raise HTTPException(422,'Định dạng xuất không hợp lệ')
@app.get('/api/ai/documents/{document_id}/file')
def ai_document_file(document_id:int,u=Depends(current),s:Session=Depends(db)):
 require_permission(s,u,'AI_CHAT','Bạn không có quyền dùng trợ lý AI')
 item=get_ai_document(document_id)
 if not item:raise HTTPException(404,'Không tìm thấy văn bản trong chỉ mục')
 path=Path(item['absolute_path']).resolve();root=AI_DOCUMENTS_ROOT.resolve()
 try:path.relative_to(root)
 except ValueError:raise HTTPException(403,'Đường dẫn văn bản không hợp lệ')
 if not path.is_file():raise HTTPException(404,'File gốc không còn tồn tại. Vui lòng đồng bộ lại chỉ mục AI')
 return FileResponse(path,filename=item['file_name'],content_disposition_type='inline')
@app.post('/api/ai/reindex')
def ai_reindex(u=Depends(current)):
 if u.role!='ADMIN':raise HTTPException(403,'Chỉ Admin được đồng bộ lại chỉ mục AI')
 try:return ai_sync_folder(str(AI_DOCUMENTS_ROOT),workers=2,use_cache=False,force_reextract=True)
 except RuntimeError as exc:raise HTTPException(409,str(exc)) from exc
 except Exception as exc:raise HTTPException(500,f'Đồng bộ thất bại: {exc}') from exc
@app.get('/api/health')
def health():return {'status':'ok'}

# Serve the compiled React application from the same process in production.
# This catch-all mount must remain after every API route.
FRONTEND_DIST=ROOT.parent/'frontend'/'dist'
if FRONTEND_DIST.is_dir():
 app.mount('/assets',StaticFiles(directory=FRONTEND_DIST/'assets'),name='frontend-assets')
 @app.get('/{full_path:path}',include_in_schema=False)
 def frontend_app(full_path:str):
  if full_path.startswith('api/'):raise HTTPException(404,'Không tìm thấy API')
  return FileResponse(FRONTEND_DIST/'index.html')
