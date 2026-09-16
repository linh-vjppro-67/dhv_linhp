from pathlib import Path
from typing import List, Optional
import json

import pymupdf
import pytesseract
from PIL import Image
from docx import Document

from .config import (
    CACHE_DIR,
    CACHE_VERSION,
    MIN_NATIVE_TEXT_CHARS,
    OCR_DPI,
    OCR_IMAGE_COVERAGE_THRESHOLD,
    OCR_IMAGE_TEXT_THRESHOLD,
    OCR_LANGUAGES,
    QUIET_TESSERACT,
)
from .normalize import normalize_display_text
from .types import ExtractedUnit
from .utils import sha256_file, suppress_native_stderr


def _cache_path(file_hash: str) -> Path:
    return CACHE_DIR / f"{file_hash}.json"


def _load_cache(file_hash: str):
    p = _cache_path(file_hash)

    if not p.exists():
        return None

    try:
        data = json.loads(p.read_text(encoding="utf-8"))

        if data.get("cache_version") != CACHE_VERSION:
            return None

        if data.get("ocr_languages") != OCR_LANGUAGES:
            return None

        if int(data.get("ocr_dpi", OCR_DPI)) != OCR_DPI:
            return None

        return [
            ExtractedUnit(
                text=item["text"],
                page_no=item.get("page_no"),
                extraction_mode=item["extraction_mode"],
            )
            for item in data.get("units", [])
        ]

    except Exception:
        return None


def _save_cache(
    file_hash: str,
    units: List[ExtractedUnit],
):
    payload = {
        "cache_version": CACHE_VERSION,
        "ocr_languages": OCR_LANGUAGES,
        "ocr_dpi": OCR_DPI,
        "units": [
            {
                "text": u.text,
                "page_no": u.page_no,
                "extraction_mode": u.extraction_mode,
            }
            for u in units
        ],
    }

    _cache_path(file_hash).write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def _largest_image_coverage(page) -> float:
    """
    Ước lượng trang có phải scan nguyên trang hay không.
    Logo nhỏ không làm trang bị OCR lại.
    """
    page_area = max(float(page.rect.width * page.rect.height), 1.0)
    best = 0.0

    try:
        for image in page.get_images(full=True):
            xref = image[0]

            for rect in page.get_image_rects(xref):
                ratio = float(rect.width * rect.height) / page_area
                best = max(best, ratio)
    except Exception:
        return 0.0

    return min(best, 1.0)


def _page_needs_ocr(page, native_text: str) -> bool:
    if len(native_text) < MIN_NATIVE_TEXT_CHARS:
        return True

    if len(native_text) < OCR_IMAGE_TEXT_THRESHOLD:
        coverage = _largest_image_coverage(page)
        if coverage >= OCR_IMAGE_COVERAGE_THRESHOLD:
            return True

    return False


def _native_text_is_scrambled(native_text: str) -> bool:
    """Detect PDFs whose embedded glyph order produces one unusable mega-line."""
    lines = [line.strip() for line in native_text.splitlines() if line.strip()]
    if len(native_text) < 300 or not lines:
        return False
    return max(map(len, lines)) > max(500, len(native_text) * 0.55)


def _ocr_page_image(page) -> str:
    scale = OCR_DPI / 72
    pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    # PSM 4 preserves administrative-document blocks much better than reading
    # the PDF's often corrupted embedded glyph order.
    return normalize_display_text(
        pytesseract.image_to_string(
            image, lang=OCR_LANGUAGES, config="--psm 4"
        )
    )


def extract_pdf(path: Path) -> List[ExtractedUnit]:
    units: List[ExtractedUnit] = []

    with pymupdf.open(path) as doc:
        for idx, page in enumerate(doc):
            page_no = idx + 1

            native_text = normalize_display_text(
                page.get_text(
                    "text",
                    sort=True,
                )
            )

            force_ocr = _native_text_is_scrambled(native_text)
            if not force_ocr and not _page_needs_ocr(page, native_text):
                if native_text:
                    units.append(
                        ExtractedUnit(
                            text=native_text,
                            page_no=page_no,
                            extraction_mode="native_pdf",
                        )
                    )
                continue

            try:
                with suppress_native_stderr(QUIET_TESSERACT):
                    ocr_text = _ocr_page_image(page)

            except Exception as exc:
                raise RuntimeError(
                    f"OCR lỗi: {path.name}, trang {page_no}. "
                    f"Kiểm tra Tesseract + language pack "
                    f"'{OCR_LANGUAGES}'. Chi tiết: {exc}"
                ) from exc

            if force_ocr or len(ocr_text) >= len(native_text):
                final_text = ocr_text
                mode = "ocr_pdf"
            else:
                final_text = native_text
                mode = "native_pdf"

            if final_text:
                units.append(
                    ExtractedUnit(
                        text=final_text,
                        page_no=page_no,
                        extraction_mode=mode,
                    )
                )

    return units


def _table_to_text(table) -> str:
    rows = []

    for row in table.rows:
        cells = [
            normalize_display_text(cell.text)
            for cell in row.cells
        ]

        if any(cells):
            rows.append(" | ".join(cells))

    return "\n".join(rows)


def extract_docx(path: Path) -> List[ExtractedUnit]:
    doc = Document(path)
    blocks = []

    if hasattr(doc, "iter_inner_content"):
        for item in doc.iter_inner_content():
            if hasattr(item, "rows"):
                text = _table_to_text(item)
            else:
                text = normalize_display_text(item.text)

            if text:
                blocks.append(text)
    else:
        for p in doc.paragraphs:
            text = normalize_display_text(p.text)
            if text:
                blocks.append(text)

        for table in doc.tables:
            text = _table_to_text(table)
            if text:
                blocks.append(text)

    text = normalize_display_text("\n\n".join(blocks))

    return (
        [
            ExtractedUnit(
                text=text,
                page_no=None,
                extraction_mode="docx",
            )
        ]
        if text
        else []
    )


def extract_text_file(path: Path) -> List[ExtractedUnit]:
    for enc in (
        "utf-8",
        "utf-8-sig",
        "cp1258",
        "latin-1",
    ):
        try:
            text = normalize_display_text(
                path.read_text(encoding=enc)
            )

            return (
                [
                    ExtractedUnit(
                        text=text,
                        page_no=None,
                        extraction_mode="txt",
                    )
                ]
                if text
                else []
            )

        except UnicodeDecodeError:
            continue

    return []


def extract_file(
    path: Path,
    file_hash: Optional[str] = None,
    use_cache: bool = True,
):
    file_hash = file_hash or sha256_file(path)

    if use_cache:
        cached = _load_cache(file_hash)

        if cached is not None:
            return file_hash, cached, True

    ext = path.suffix.lower()

    if ext == ".pdf":
        units = extract_pdf(path)
    elif ext == ".docx":
        units = extract_docx(path)
    elif ext in {".txt", ".md"}:
        units = extract_text_file(path)
    else:
        raise ValueError(
            f"Không hỗ trợ file: {path.suffix}"
        )

    _save_cache(file_hash, units)

    return file_hash, units, False
